"""Manual-check workbook: the tool's closed-form numbers recomputed by hand, as live Excel formulas, beside the tool's own.

For one analysis case this writes an .xlsx in which every quantity that has a closed form is recomputed with ordinary
spreadsheet formulas from the entered inputs (Beta-PERT / triangular / uniform / fixed moments, event EMV, the 5 x 5
class and level, the net benefit of each response, expected total delay and probability of any delay) and compared with
the value the tool printed.  Each comparison row carries a difference and a PASS / FAIL formula.

What it does NOT recompute: the simulated cost and duration percentiles, the option comparison, and anything that needs
the full distribution (liquidated damages, P90).  Those are checked by the validation scripts in
report/scripts/validation_2026-10-09/.  The workbook states this on its first sheet.

    data = manual_check_xlsx(case)                  bytes of the .xlsx
    data = manual_check_xlsx(case, result=res)      reuse a run_analysis() result for the same case

Excel formulas written here are recalculated when the file is opened (cached values are absent until then).
"""
from __future__ import annotations

import io
from typing import Mapping, Optional

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
_BLUE = Font(name=FONT, size=10, color="0000FF")
_BLK = Font(name=FONT, size=10)
_BOLD = Font(name=FONT, size=10, bold=True)
_GRN = Font(name=FONT, size=10, color="008000")
_TITLE = Font(name=FONT, size=14, bold=True)
_HF = Font(name=FONT, size=10, bold=True, color="FFFFFF")
_HFILL = PatternFill("solid", start_color="1F3864")
_INP = PatternFill("solid", start_color="FFF2CC")
_THIN = Side(style="thin", color="BFBFBF")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP = Alignment(wrap_text=True, vertical="top")
_CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
TOL = "0.000001"          # absolute tolerance for closed-form comparisons (relative part added in the formulas)


def _put(ws, ref, v, font=_BLK, fmt=None, fill=None, al=None):
    c = ws[ref]
    c.value = v
    c.font = font
    c.border = _BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if al:
        c.alignment = al
    return c


def _head(ws, row, labels, height=34):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=1 + i, value=t)
        c.font, c.fill, c.alignment, c.border = _HF, _HFILL, _CTR, _BOX
    ws.row_dimensions[row].height = height


def _widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def _pf(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], font=Font(name=FONT, color="006100", bold=True),
                                                  fill=PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], font=Font(name=FONT, color="9C0006", bold=True),
                                                  fill=PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")))


def _mean_f(kind, a, m, b, lam):
    return f'IF({kind}="pert",({a}+{lam}*{m}+{b})/({lam}+2),IF({kind}="triangular",({a}+{m}+{b})/3,IF({kind}="uniform",({a}+{b})/2,{m})))'


def _var_f(kind, a, m, b, lam):
    mu = f"({_mean_f(kind, a, m, b, lam)})"
    return (f'IF({kind}="pert",({mu}-{a})*({b}-{mu})/({lam}+3),IF({kind}="triangular",({a}^2+{m}^2+{b}^2-{a}*{m}-{a}*{b}-{m}*{b})/18,'
            f'IF({kind}="uniform",({b}-{a})^2/12,0)))')


def _num(x, default=None):
    try:
        return float(x)
    except (TypeError, ValueError):
        return default


def manual_check_xlsx(case: Mapping, result: Optional[Mapping] = None, evidence_index: Optional[Mapping[str, str]] = None) -> bytes:
    """Build the workbook for `case`.  Raises service.CaseError for an invalid case (same as run_analysis)."""
    from app import analysis, service

    pa = analysis.parse_analysis(case)                       # raises CaseError listing every problem
    res = result if result is not None else service.run_analysis(case, evidence_index)
    mx = {m["risk_id"]: m for m in service.run_case(case, evidence_index)["matrix"]}
    ms = service.run_case(case, evidence_index)["matrix_settings"]
    risks, mits, cost = pa["risks"], pa["mitigations"], pa["cost"]
    emv = {r["risk_id"]: r for r in res.get("event_emv") or []}
    chk = {c["mitigation_id"]: c for c in res.get("mitigation_checks") or []}
    summ = res.get("summary") or {}
    corr = bool(pa.get("correlation"))
    k = len(risks)

    wb = Workbook()
    # ------------------------------------------------------------------------------------------------ README
    ws = wb.active
    ws.title = "README"
    ws["A1"].value, ws["A1"].font = "Manual check of the tool's calculations", _TITLE
    name = (res.get("project") or {}).get("name") or "case"
    rows = [
        ("Case", f"{name}. Activity: {(res.get('activity') or {}).get('name', '')}."),
        ("What this is", "Closed-form quantities recomputed by hand as live Excel formulas from the inputs on the Inputs sheet, beside the number the tool printed, with a difference and a PASS / FAIL formula. If you change a blue input, the hand column recomputes but the pasted tool column does not, so differences will then appear."),
        ("Checked here", "Delay moments by distribution kind; event EMV = p x (mean delay x delay cost + direct cost); 5 x 5 class and level; net benefit of each response (EMV removed - secondary EMV - cost); expected total delay and probability of any delay against the simulation (within 3 standard errors)."),
        ("NOT checked here", "Simulated duration and cost percentiles (P50, P90), liquidated-damage costs, and the option comparison need the full distribution. They are checked by the validation scripts in report/scripts/validation_2026-10-09/ (exact grid convolution and a large independent simulation)."),
        ("Occurrence correlation", "A correlation between risks was entered. The probability of any delay assumes independence, so that row is marked n/a." if corr else "No occurrence correlation was entered; the independence formulas apply."),
        ("Colour legend", "Blue on yellow = input. Blue on white = value pasted from the tool. Black = formula. Green = link to another sheet."),
        ("Not evidence", "A PASS shows the arithmetic is right, not that the entered probabilities, delays and costs describe a real site."),
        ("Tool run", f"n = {res.get('n_used')}, seed = {res.get('seed')}, criterion = {res.get('criterion')}."),
    ]
    for i, (a, b) in enumerate(rows):
        r = 3 + i
        _put(ws, f"A{r}", a, _BOLD, al=_WRAP)
        _put(ws, f"B{r}", b, _BLK, al=_WRAP)
        ws.row_dimensions[r].height = 48
    ov = 3 + len(rows) + 1
    _put(ws, f"A{ov}", "Overall", _BOLD)
    _widths(ws, [22, 120])

    # ------------------------------------------------------------------------------------------------ Inputs
    wi = wb.create_sheet("Inputs")
    wi["A1"].value, wi["A1"].font = "Inputs (from the case)", _TITLE
    _head(wi, 3, ["Quantity", "Value", "Unit"], 22)
    dl = pa["deadline"].value if pa["deadline"] else None
    scal = [("Planned duration T0", pa["planned"].value, "days"),
            ("Deadline", dl, "days (blank = none)"),
            ("Delay cost per day Cd", cost.cost_per_delay_day.value, f"{cost.currency} per day"),
            ("Liquidated damages per day", cost.ld_per_day_after_deadline.value, f"{cost.currency} per day (not used in these checks)"),
            ("Simulation draws n", res.get("n_used"), "draws"),
            ("Seed", res.get("seed"), "")]
    for i, (a, v, u) in enumerate(scal):
        r = 4 + i
        _put(wi, f"A{r}", a, _BOLD)
        _put(wi, f"B{r}", v, _BLUE, "#,##0.####", _INP)
        _put(wi, f"C{r}", u)
    T0, CD, NSIM = "Inputs!$B$4", "Inputs!$B$6", "Inputs!$B$8"
    rh = 12
    _head(wi, rh, ["Risk id", "Name", "Delay kind", "a (min)", "m (mode)", "b (max)", "lambda", "p", "Direct cost if it occurs"], 30)
    for i, r in enumerate(risks):
        rr = rh + 1 + i
        d = r.delay
        _put(wi, f"A{rr}", r.risk_id, _BOLD)
        _put(wi, f"B{rr}", r.name)
        _put(wi, f"C{rr}", d.kind, _BLUE, None, _INP)
        for col, v in zip("DEFG", (d.a, d.m, d.b, d.lam)):
            _put(wi, f"{col}{rr}", v, _BLUE, "0.0###", _INP)
        _put(wi, f"H{rr}", r.p.value, _BLUE, "0.0###", _INP)
        _put(wi, f"I{rr}", r.direct_cost_if_occurs.value, _BLUE, "#,##0", _INP)
    R0, R1 = rh + 1, rh + k
    mh = R1 + 3
    _head(wi, mh, ["Response id", "Targets risk", "Cost", "p after", "Delay-after kind", "a", "m", "b", "lambda", "Feasible?"], 30)
    for i, m in enumerate(mits):
        rr = mh + 1 + i
        da = m.delay_after
        _put(wi, f"A{rr}", m.mitigation_id, _BOLD)
        _put(wi, f"B{rr}", m.risk_id)
        _put(wi, f"C{rr}", m.cost.value, _BLUE, "#,##0", _INP)
        _put(wi, f"D{rr}", m.p_after.value if m.p_after else None, _BLUE, "0.0###", _INP)
        _put(wi, f"E{rr}", da.kind if da else None, _BLUE, None, _INP)
        for col, v in zip("FGHI", (da.a, da.m, da.b, da.lam) if da else (None,) * 4):
            _put(wi, f"{col}{rr}", v, _BLUE, "0.0###", _INP)
        _put(wi, f"J{rr}", "Yes" if m.feasible else "No", _BLUE, None, _INP)
    M0, M1 = mh + 1, mh + max(1, len(mits))
    sh = M1 + 3
    _head(wi, sh, ["Response id", "Secondary risk id", "p", "Delay kind", "a", "m", "b", "lambda", "Direct cost"], 30)
    sec_rows = [(m.mitigation_id, s) for m in mits for s in m.secondary_risks]
    for i, (mid, s) in enumerate(sec_rows):
        rr = sh + 1 + i
        _put(wi, f"A{rr}", mid, _BOLD)
        _put(wi, f"B{rr}", s.risk_id)
        _put(wi, f"C{rr}", s.p.value, _BLUE, "0.0###", _INP)
        _put(wi, f"D{rr}", s.delay.kind, _BLUE, None, _INP)
        for col, v in zip("EFGH", (s.delay.a, s.delay.m, s.delay.b, s.delay.lam)):
            _put(wi, f"{col}{rr}", v, _BLUE, "0.0###", _INP)
        _put(wi, f"I{rr}", s.direct_cost_if_occurs.value, _BLUE, "#,##0", _INP)
    S0, S1 = sh + 1, sh + max(1, len(sec_rows))
    eh = S1 + 3
    _put(wi, f"A{eh}", "Probability class edges (lower-edge inclusive)", _BOLD)
    _put(wi, f"A{eh+1}", "Impact edges (fraction of T0)", _BOLD)
    _put(wi, f"A{eh+2}", "Level cut-offs on score (Low <=, Moderate <=, High <=)", _BOLD)
    for j, v in enumerate(ms["p_edges"]):
        _put(wi, f"{get_column_letter(2+j)}{eh}", v, _BLUE, "0.0###", _INP)
    for j, v in enumerate(ms["impact_edges"]):
        _put(wi, f"{get_column_letter(2+j)}{eh+1}", v, _BLUE, "0.0###", _INP)
    for j, v in enumerate(ms["level_thresholds"]):
        _put(wi, f"{get_column_letter(2+j)}{eh+2}", v, _BLUE, "0", _INP)
    _widths(wi, [44, 34, 14, 12, 14, 12, 12, 12, 22, 12])

    # ------------------------------------------------------------------------------------------------ Risks
    wr = wb.create_sheet("Risks")
    wr["A1"].value, wr["A1"].font = "Per-risk moments and event EMV: hand vs tool", _TITLE
    _head(wr, 3, ["Risk id", "p", "Mean delay if it occurs (d)", "Variance (d^2)", "Conditional cost\n=mean x Cd + direct", "Event EMV\n=p x conditional cost",
                  "Tool mean (d)", "Tool EMV", "Mean diff", "EMV diff", "Result"], 52)
    for i, r in enumerate(risks):
        rr, ir = 4 + i, R0 + i
        kd, a, m, b, lam = (f"Inputs!$C{ir}", f"Inputs!$D{ir}", f"Inputs!$E{ir}", f"Inputs!$F{ir}", f"Inputs!$G{ir}")
        _put(wr, f"A{rr}", f"=Inputs!A{ir}", _GRN)
        _put(wr, f"B{rr}", f"=Inputs!H{ir}", _GRN, "0.0###")
        _put(wr, f"C{rr}", "=" + _mean_f(kd, a, m, b, lam), _BLK, "0.0000")
        _put(wr, f"D{rr}", "=" + _var_f(kd, a, m, b, lam), _BLK, "0.000000")
        _put(wr, f"E{rr}", f"=C{rr}*{CD}+Inputs!I{ir}", _BLK, "#,##0.00")
        _put(wr, f"F{rr}", f"=B{rr}*E{rr}", _BLK, "#,##0.00")
        t = emv.get(r.risk_id, {})
        _put(wr, f"G{rr}", t.get("expected_delay_if_occurs_days"), _BLUE, "0.0000")
        _put(wr, f"H{rr}", t.get("emv"), _BLUE, "#,##0.00")
        _put(wr, f"I{rr}", f"=C{rr}-G{rr}", _BLK, "0.0E+00")
        _put(wr, f"J{rr}", f"=F{rr}-H{rr}", _BLK, "0.0E+00")
        _put(wr, f"K{rr}", f'=IF(AND(ABS(I{rr})<={TOL}*(1+ABS(G{rr})),ABS(J{rr})<={TOL}*(1+ABS(H{rr}))),"PASS","FAIL")', _BOLD, al=_CTR)
    K0, K1 = 4, 3 + k
    _pf(wr, f"K{K0}:K{K1}")
    _widths(wr, [14, 8, 20, 16, 22, 18, 12, 14, 12, 12, 10])

    # ------------------------------------------------------------------------------------------------ Totals
    wt = wb.create_sheet("Totals")
    wt["A1"].value, wt["A1"].font = "Expected total delay and probability of any delay: hand vs simulation", _TITLE
    _head(wt, 3, ["Quantity", "Hand (formula)", "Tool (simulated)", "Difference", "Tolerance (3 standard errors)", "Result"], 44)
    se_mean = _num(summ.get("mean_se"), 0.0)
    # 1 - product(1-p): helper column on the Risks sheet keeps the formula simple and portable
    _put(wr, "M3", "1 - p", _HF, None, _HFILL, _CTR)
    for i in range(k):
        _put(wr, f"M{4+i}", f"=1-B{4+i}", _BLK, "0.0###")
    _put(wt, "A4", "Expected total delay (d) = sum of p x mean", _BOLD)
    _put(wt, "B4", f"=SUMPRODUCT(Risks!$B${K0}:$B${K1},Risks!$C${K0}:$C${K1})", _BLK, "0.0000")
    _put(wt, "C4", summ.get("expected_delay"), _BLUE, "0.0000")
    _put(wt, "D4", "=B4-C4", _BLK, "0.0000")
    _put(wt, "E4", 3 * se_mean, _BLUE, "0.0000")
    _put(wt, "F4", '=IF(ABS(D4)<=E4,"PASS","FAIL")', _BOLD, al=_CTR)
    _put(wt, "A5", "Probability of any delay = 1 - product(1 - p)", _BOLD)
    _put(wt, "B5", f"=1-PRODUCT(Risks!$M${K0}:$M${K1})", _BLK, "0.0000")
    _put(wt, "C5", summ.get("p_any_delay"), _BLUE, "0.0000")
    _put(wt, "D5", "=B5-C5", _BLK, "0.0000")
    pa_hand = 1.0
    for r in risks:
        pa_hand *= 1.0 - r.p.value
    pa_hand = 1.0 - pa_hand
    n_used = max(1, int(res.get("n_used") or 1))
    _put(wt, "E5", 3 * (max(pa_hand * (1 - pa_hand), 1e-12) / n_used) ** 0.5, _BLUE, "0.0000")
    _put(wt, "F5", ('n/a' if corr else '=IF(ABS(D5)<=E5,"PASS","FAIL")'), _BOLD, al=_CTR)
    _put(wt, "A6", "Sum of event EMV (cost units)", _BOLD)
    _put(wt, "B6", f"=SUM(Risks!$F${K0}:$F${K1})", _BLK, "#,##0.00")
    _put(wt, "C6", sum((emv[r.risk_id]["emv"] for r in risks if r.risk_id in emv)), _BLUE, "#,##0.00")
    _put(wt, "D6", "=B6-C6", _BLK, "0.0E+00")
    _put(wt, "E6", f"={TOL}*(1+ABS(C6))", _BLK, "0.0E+00")
    _put(wt, "F6", '=IF(ABS(D6)<=E6,"PASS","FAIL")', _BOLD, al=_CTR)
    _put(wt, "A8", "Tolerance for the first two rows uses the tool's own standard error of the mean and the binomial standard error at the number of draws used. A simulation will miss its exact value by more than 3 standard errors about 3 times in 1,000.", _BLK, al=_WRAP)
    wt.merge_cells("A8:F8")
    wt.row_dimensions[8].height = 44
    _pf(wt, "F4:F6")
    _widths(wt, [48, 18, 18, 14, 22, 10])

    # ------------------------------------------------------------------------------------------------ Matrix
    wm = wb.create_sheet("Matrix")
    wm["A1"].value, wm["A1"].font = "5 x 5 class and level: hand vs tool", _TITLE
    _head(wm, 3, ["Risk id", "p", "Mean delay (d)", "Ratio = mean / T0", "Probability class", "Impact class", "Score", "Level", "Tool p class", "Tool impact class", "Tool score", "Tool level", "Result"], 40)
    pe = [f"Inputs!${get_column_letter(2+j)}${eh}" for j in range(len(ms["p_edges"]))]
    ie = [f"Inputs!${get_column_letter(2+j)}${eh+1}" for j in range(len(ms["impact_edges"]))]
    lv = [f"Inputs!${get_column_letter(2+j)}${eh+2}" for j in range(len(ms["level_thresholds"]))]
    tol_e = "0.000000000001"
    for i, r in enumerate(risks):
        rr = 4 + i
        _put(wm, f"A{rr}", f"=Risks!A{4+i}", _GRN)
        _put(wm, f"B{rr}", f"=Risks!B{4+i}", _GRN, "0.0###")
        _put(wm, f"C{rr}", f"=Risks!C{4+i}", _GRN, "0.0000")
        _put(wm, f"D{rr}", f"=C{rr}/{T0}", _BLK, "0.00000")
        _put(wm, f"E{rr}", "=1" + "".join(f"+(B{rr}>={e}-{tol_e})" for e in pe), _BLK, "0")
        _put(wm, f"F{rr}", "=1" + "".join(f"+(D{rr}>={e}-{tol_e})" for e in ie), _BLK, "0")
        _put(wm, f"G{rr}", f"=E{rr}*F{rr}", _BLK, "0")
        _put(wm, f"H{rr}", f'=IF(G{rr}<={lv[0]},"Low",IF(G{rr}<={lv[1]},"Moderate",IF(G{rr}<={lv[2]},"High","Extreme")))', _BLK)
        t = mx.get(r.risk_id, {})
        _put(wm, f"I{rr}", t.get("p_class"), _BLUE, "0")
        _put(wm, f"J{rr}", t.get("impact_class"), _BLUE, "0")
        _put(wm, f"K{rr}", t.get("score"), _BLUE, "0")
        _put(wm, f"L{rr}", t.get("level"), _BLUE)
        _put(wm, f"M{rr}", f'=IF(AND(E{rr}=I{rr},F{rr}=J{rr},G{rr}=K{rr},H{rr}=L{rr}),"PASS","FAIL")', _BOLD, al=_CTR)
    _pf(wm, f"M4:M{3+k}")
    _widths(wm, [14, 8, 16, 16, 18, 14, 9, 12, 12, 14, 11, 12, 10])

    # ------------------------------------------------------------------------------------------------ Responses
    wp = wb.create_sheet("Responses")
    wp["A1"].value, wp["A1"].font = "Net benefit of each response: hand vs tool (this response alone, analytic, linear cost)", _TITLE
    nm = max(1, len(mits))
    sec_start = 6 + nm + 3                                  # first row of the secondary block
    _head(wp, 3, ["Response", "Targets", "EMV before", "p after (or original p)", "Mean delay after (d)", "Direct cost", "EMV after", "Secondary EMV", "Cost", "Net benefit\n=before - after - secondary - cost",
                  "Break-even: cost / EMV before", "Feasible?", "Tool EMV before", "Tool EMV after", "Tool secondary", "Tool net", "Net diff", "Result"], 64)
    for i, m in enumerate(mits):
        rr, ir = 4 + i, M0 + i
        rk = f"MATCH(Inputs!$B{ir},Risks!$A${K0}:$A${K1},0)"
        rin = f"MATCH(Inputs!$B{ir},Inputs!$A${R0}:$A${R1},0)"
        _put(wp, f"A{rr}", f"=Inputs!A{ir}", _GRN)
        _put(wp, f"B{rr}", f"=Inputs!B{ir}", _GRN)
        _put(wp, f"C{rr}", f"=INDEX(Risks!$F${K0}:$F${K1},{rk})", _GRN, "#,##0.00")
        _put(wp, f"D{rr}", f"=IF(ISNUMBER(Inputs!D{ir}),Inputs!D{ir},INDEX(Risks!$B${K0}:$B${K1},{rk}))", _BLK, "0.0###")
        _put(wp, f"E{rr}", "=IF(ISNUMBER(Inputs!G" + str(ir) + "),"
                           + _mean_f(f"Inputs!$E{ir}", f"Inputs!$F{ir}", f"Inputs!$G{ir}", f"Inputs!$H{ir}", f"Inputs!$I{ir}")
                           + f",INDEX(Risks!$C${K0}:$C${K1},{rk}))", _BLK, "0.0000")
        _put(wp, f"F{rr}", f"=INDEX(Inputs!$I${R0}:$I${R1},{rin})", _GRN, "#,##0.00")
        _put(wp, f"G{rr}", f"=D{rr}*(E{rr}*{CD}+F{rr})", _BLK, "#,##0.00")
        _put(wp, f"H{rr}", f"=SUMIFS($E${sec_start}:$E${sec_start+len(sec_rows) if sec_rows else sec_start},$A${sec_start}:$A${sec_start+len(sec_rows) if sec_rows else sec_start},A{rr})", _BLK, "#,##0.00")
        _put(wp, f"I{rr}", f"=Inputs!C{ir}", _GRN, "#,##0")
        _put(wp, f"J{rr}", f"=C{rr}-G{rr}-H{rr}-I{rr}", _BLK, "#,##0.00;(#,##0.00);-")
        _put(wp, f"K{rr}", f'=IF(C{rr}=0,"n/a",I{rr}/C{rr})', _BLK, "0.0000")
        _put(wp, f"L{rr}", f"=Inputs!J{ir}", _GRN, al=_CTR)
        t = chk.get(m.mitigation_id, {})
        _put(wp, f"M{rr}", t.get("emv_before"), _BLUE, "#,##0.00")
        _put(wp, f"N{rr}", t.get("emv_after"), _BLUE, "#,##0.00")
        _put(wp, f"O{rr}", t.get("secondary_risk_emv"), _BLUE, "#,##0.00")
        _put(wp, f"P{rr}", t.get("net_benefit"), _BLUE, "#,##0.00;(#,##0.00);-")
        _put(wp, f"Q{rr}", f"=J{rr}-P{rr}", _BLK, "0.0E+00")
        _put(wp, f"R{rr}", f'=IF(AND(ABS(C{rr}-M{rr})<={TOL}*(1+ABS(M{rr})),ABS(G{rr}-N{rr})<={TOL}*(1+ABS(N{rr})),ABS(H{rr}-O{rr})<={TOL}*(1+ABS(O{rr})),ABS(Q{rr})<={TOL}*(1+ABS(P{rr}))),"PASS","FAIL")', _BOLD, al=_CTR)
    if mits:
        _pf(wp, f"R4:R{3+len(mits)}")
    _put(wp, f"A{sec_start-2}", "Secondary risks created by a response (hand EMV = p x (mean x Cd + direct cost))", _BOLD)
    _head(wp, sec_start - 1, ["Response", "Secondary risk", "p", "Mean delay (d)", "EMV"], 28)
    for i, (mid, s) in enumerate(sec_rows):
        rr, ir = sec_start + i, S0 + i
        _put(wp, f"A{rr}", f"=Inputs!A{ir}", _GRN)
        _put(wp, f"B{rr}", f"=Inputs!B{ir}", _GRN)
        _put(wp, f"C{rr}", f"=Inputs!C{ir}", _GRN, "0.0###")
        _put(wp, f"D{rr}", "=" + _mean_f(f"Inputs!$D{ir}", f"Inputs!$E{ir}", f"Inputs!$F{ir}", f"Inputs!$G{ir}", f"Inputs!$H{ir}"), _BLK, "0.0000")
        _put(wp, f"E{rr}", f"=C{rr}*(D{rr}*{CD}+Inputs!I{ir})", _BLK, "#,##0.00")
    _widths(wp, [12, 12, 14, 14, 16, 12, 14, 14, 10, 26, 18, 10, 14, 14, 14, 14, 12, 10])

    # ------------------------------------------------------------------------------------------------ Overall
    ref = []
    ref.append(f"Risks!K{K0}:K{K1}")
    ref.append("Totals!F4:F6")
    ref.append(f"Matrix!M4:M{3+k}")
    if mits:
        ref.append(f"Responses!R4:R{3+len(mits)}")
    p_ = "+".join(f'COUNTIF({r},"PASS")' for r in ref)
    f_ = "+".join(f'COUNTIF({r},"FAIL")' for r in ref)
    _put(ws, f"B{ov}", f'=IF(({f_})=0,"ALL "&({p_})&" CHECKS PASS",({f_})&" OF "&(({p_})+({f_}))&" CHECKS FAIL")', _BOLD)
    for s in wb.worksheets:
        s.sheet_view.showGridLines = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
