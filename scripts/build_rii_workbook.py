"""Build RII_Masonry_Literature.xlsx from data/literature_seed.json (the same file the app's literature seed reads).

Sheets: How it works · Studies · RII_Data (every value copied from the papers) · Seed_Calculation (live formulas that
reproduce the app's seed) · Validation (Spearman rank correlation against held-out studies) · Other_indices.

Run: python scripts/build_rii_workbook.py [output.xlsx]
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
DATA = json.loads((ROOT / "data" / "literature_seed.json").read_text(encoding="utf-8-sig"))
LIB = json.loads((ROOT / "data" / "risk_library_masonry.json").read_text(encoding="utf-8-sig"))
RISKS = [r["id"] for r in LIB]
NAMES = {r["id"]: r["name"] for r in LIB}
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "RII_Masonry_Literature.xlsx"

FONT = "Arial"
BLUE, GREEN, BLACK, GREY = "0000FF", "008000", "000000", "6B6B62"
HEAD_FILL = PatternFill("solid", fgColor="2B2B26")
SUB_FILL = PatternFill("solid", fgColor="E9E9DF")
ASSUME_FILL = PatternFill("solid", fgColor="FFF2A8")
THIN = Side(style="thin", color="C8C8BE")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def f(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def head(ws, row, values, widths=None):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = f("FFFFFF", bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=row, column=c).column_letter].width = w


def put(ws, row, col, value, color=BLACK, bold=False, fmt=None, fill=None, wrap=False, italic=False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = f(color, bold, italic=italic)
    cell.border = BOX
    cell.alignment = Alignment(wrap_text=wrap, vertical="top")
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = f(bold=True, size=14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(GREY, italic=True)


wb = Workbook()

# ------------------------------------------------------------------ Studies
st = wb.active
st.title = "Studies"
title(st, "Studies the RII values come from", "Role and the mapping column on RII_Data (yellow) are judgements you may change; everything else is copied from the papers.")
cols = ["ID", "Authors", "Year", "Title", "Venue", "DOI", "Country", "Context", "Masonry-specific?", "Respondents (N)",
        "Rating scale", "Where in the paper", "How the values were obtained", "Role", "Note"]
head(st, 4, cols, [8, 26, 6, 48, 30, 30, 22, 34, 10, 12, 10, 18, 38, 11, 48])
for i, s in enumerate(DATA["studies"], 5):
    vals = [s["id"], s["authors"], s["year"], s["title"], s["venue"], s["doi"], s["country"], s["context"], s["masonry"],
            s["n"], s["scale"], s["location"], s["obtained"], s["role"].capitalize(), s.get("note", "")]
    for c, v in enumerate(vals, 1):
        put(st, i, c, v, BLUE if c != 14 else BLUE, wrap=c in (2, 4, 5, 7, 8, 13, 15), fill=ASSUME_FILL if c == 14 else None)
ST_LAST = 4 + len(DATA["studies"])
st.freeze_panes = "B5"

# ------------------------------------------------------------------ RII_Data
rd = wb.create_sheet("RII_Data")
title(rd, "Every RII value recorded, one row per factor as printed in the paper",
      "Blue = copied from the paper. Green = looked up from Studies. Black = formula. Yellow = mapping judgement (Direct / Related / blank).")
head(rd, 4, ["Study ID", "Country", "Role", "Factor (as printed)", "RII (0-1)", "Rank in study", "Tool risk", "Mapping",
             "Used by the model?", "Note", "z-score within study"], [9, 24, 11, 52, 10, 9, 10, 10, 11, 52, 12])
rows = DATA["rows"]
for i, r in enumerate(rows, 5):
    put(rd, i, 1, r["study"], BLUE)
    put(rd, i, 2, f'=INDEX(Studies!$G$5:$G${ST_LAST},MATCH(A{i},Studies!$A$5:$A${ST_LAST},0))', GREEN)
    put(rd, i, 3, f'=INDEX(Studies!$N$5:$N${ST_LAST},MATCH(A{i},Studies!$A$5:$A${ST_LAST},0))', GREEN)
    put(rd, i, 4, r["factor"], BLUE, wrap=True)
    put(rd, i, 5, r["rii"], BLUE, fmt="0.000")
    put(rd, i, 6, r.get("rank"), BLUE)
    put(rd, i, 7, r.get("risk") or "", BLUE)
    put(rd, i, 8, (r.get("mapping") or "").capitalize(), BLUE, fill=ASSUME_FILL)
    put(rd, i, 9, f'=IF(AND(C{i}="Seed",H{i}="Direct",G{i}<>""),"Yes","No")')
    put(rd, i, 10, r.get("note", ""), GREY, wrap=True)
    put(rd, i, 11, f'=IF(OR(INDEX(Standardised!$B$5:$B$15,MATCH(A{i},Standardised!$A$5:$A$15,0))="Complete",INDEX(Standardised!$B$5:$B$15,MATCH(A{i},Standardised!$A$5:$A$15,0))="Truncated"),'
                    f'(E{i}-INDEX(Standardised!$D$5:$D$15,MATCH(A{i},Standardised!$A$5:$A$15,0)))/INDEX(Standardised!$E$5:$E$15,MATCH(A{i},Standardised!$A$5:$A$15,0)),"")',
        GREEN, fmt="0.00")
RD_FIRST, RD_LAST = 5, 4 + len(rows)
rd.freeze_panes = "E5"
rd.auto_filter.ref = f"A4:K{RD_LAST}"
rd.conditional_formatting.add(f"I{RD_FIRST}:I{RD_LAST}", CellIsRule(operator="equal", formula=['"Yes"'],
                              fill=PatternFill("solid", fgColor="CDEFC4"), font=Font(name=FONT, bold=True, color="1E5E1E")))
rng = lambda col: f"RII_Data!${col}${RD_FIRST}:${col}${RD_LAST}"   # noqa: E731

# ------------------------------------------------------------------ Seed_Calculation
sc = wb.create_sheet("Seed_Calculation")
seed_ids = [s["id"] for s in DATA["studies"] if s["role"] == "seed"]
title(sc, "Literature seed: how the model turns RII values into a place on the risk matrix",
      "Each study cell = average of that study's DIRECT values for the risk. Seed mean = average across studies. Rank -> five equal bands -> class.")
hdr = ["Risk ID", "Risk"] + seed_ids + ["Studies with a value", "Seed mean RII", "Rank", "Class (1-5)", "Matrix cell (p × impact)", "Level"]
head(sc, 4, hdr, [9, 44] + [9] * len(seed_ids) + [10, 11, 7, 9, 14, 11])
c0 = 3                                   # first study column
cN = c0 + len(seed_ids) - 1
L = lambda c: sc.cell(row=4, column=c).column_letter   # noqa: E731
first, last = 5, 4 + len(RISKS)
nrow = last + 2
for i, rid in enumerate(RISKS, first):
    put(sc, i, 1, rid, bold=True)
    put(sc, i, 2, NAMES[rid], wrap=True)
    for c in range(c0, cN + 1):
        put(sc, i, c, f'=IFERROR(AVERAGEIFS({rng("E")},{rng("A")},{L(c)}$4,{rng("G")},$A{i},{rng("H")},"Direct",{rng("C")},"Seed"),"")',
            GREEN, fmt="0.000")
    cnt, mean_, rank_, cls_, cell_, lvl_ = (L(cN + k) for k in range(1, 7))
    put(sc, i, cN + 1, f"=COUNT({L(c0)}{i}:{L(cN)}{i})")
    put(sc, i, cN + 2, f'=IF({cnt}{i}=0,"",AVERAGE({L(c0)}{i}:{L(cN)}{i}))', fmt="0.0000", bold=True)
    put(sc, i, cN + 3, f'=IF({mean_}{i}="","",RANK({mean_}{i},${mean_}${first}:${mean_}${last},0)+COUNTIFS(${mean_}${first}:${mean_}${last},{mean_}{i},$A${first}:$A${last},"<"&$A{i}))')
    put(sc, i, cN + 4, f'=IF({rank_}{i}="","",5-INT(({rank_}{i}-1)*5/$C${nrow}))', bold=True)
    put(sc, i, cN + 5, f'=IF({cls_}{i}="","not seeded","p"&{cls_}{i}&" × i"&{cls_}{i})')
    put(sc, i, cN + 6, f'=IF({cls_}{i}="","",IF({cls_}{i}^2<=5,"Low",IF({cls_}{i}^2<=10,"Moderate",IF({cls_}{i}^2<=15,"High","Extreme"))))')
MEAN_COL = L(cN + 2)
stats = [("Risks seeded (n)", f"=COUNT({MEAN_COL}{first}:{MEAN_COL}{last})", "0"),
         ("Highest mean RII", f"=MAX({MEAN_COL}{first}:{MEAN_COL}{last})", "0.0000"),
         ("Lowest mean RII", f"=MIN({MEAN_COL}{first}:{MEAN_COL}{last})", "0.0000"),
         ("Spread (highest - lowest)", f"=C{nrow + 1}-C{nrow + 2}", "0.0000"),
         ("Spread without the lowest risk", f'=C{nrow + 1}-SMALL({MEAN_COL}{first}:{MEAN_COL}{last},2)', "0.0000")]
for k, (lab, formula, fmt) in enumerate(stats):
    put(sc, nrow + k, 2, lab, bold=True)
    put(sc, nrow + k, 3, formula, fmt=fmt)
notes = [
    "How to read this sheet:",
    "1. Study columns: AVERAGEIFS over RII_Data rows with that study, that risk, mapping = Direct and role = Seed. Two direct items in one study are averaged first, so every study counts once.",
    "2. Seed mean RII = plain average of the study columns that have a value.",
    "3. Rank 1 = highest mean. Class = 5 - INT((rank - 1) x 5 / n), n = number of seeded risks (cell below). The ranks are cut into five bands of about n/5 risks each; with n = 30, ranks 1-6 -> class 5, 7-12 -> 4, 13-18 -> 3, 19-24 -> 2, 25-30 -> 1. Risks with no seed-study value stay unranked.",
    "4. The class is used for BOTH axes because an RII is an importance rating, not a probability and not days of delay. Level uses the app's thresholds: score <= 5 Low, <= 10 Moderate, <= 15 High, else Extreme.",
    "5. In the app, a site-fact rule flag raises the probability class by one (max 5), and entering your own probability and delay replaces the seed completely.",
    "6. Caution: most mean RIIs lie within a few hundredths of each other (see the spread rows). The bands make small differences look large, so treat the seed as a weak starting order.",
    "7. Ties: equal means are ordered by risk ID, the same rule the app uses (R-MTH and R-STO tie at 0.700).",
]
for k, t in enumerate(notes):
    c = sc.cell(row=nrow + 7 + k, column=1, value=t)
    c.font = f(bold=k == 0, color=BLACK if k == 0 else GREY)
sc.freeze_panes = "C5"
chart = BarChart()
chart.type = "bar"
chart.title = "Seed mean RII by risk"
chart.y_axis.title = "Mean RII (0-1)"
chart.y_axis.scaling.min = 0.5
chart.y_axis.scaling.max = 0.85
chart.add_data(Reference(sc, min_col=cN + 2, min_row=4, max_row=last), titles_from_data=True)
chart.set_categories(Reference(sc, min_col=1, min_row=first, max_row=last))
chart.legend = None
chart.varyColors = False
chart.x_axis.delete = False                       # keep the risk labels and the value axis visible
chart.y_axis.delete = False
chart.y_axis.number_format = "0.00"
chart.series[0].graphicalProperties.solidFill = "4A4A42"
chart.series[0].graphicalProperties.line.solidFill = "4A4A42"
chart.height, chart.width = 18, 16
sc.add_chart(chart, f"{L(cN + 7)}4")

SEED_MEAN = f"Seed_Calculation!${MEAN_COL}${first}:${MEAN_COL}${last}"
SEED_IDS = f"Seed_Calculation!$A${first}:$A${last}"

# ------------------------------------------------------------------ Validation
va = wb.create_sheet("Validation")
title(va, "Validation: does the seed ranking agree with studies the model did NOT use?",
      "Spearman rank correlation between the seed mean RII and each held-out study, over the risks both cover (direct items only).")
from app.literature_seed import seed_table  # noqa: E402
seeded = {k for k, v in seed_table().items() if v['rank']}
val_ids = [s["id"] for s in DATA["studies"] if s["role"] == "validation"]
direct = lambda sid: sorted({r["risk"] for r in rows if r["study"] == sid and r.get("mapping") == "direct" and r.get("risk")})  # noqa: E731
blocks = [(sid, [x for x in RISKS if x in set(direct(sid)) & seeded]) for sid in val_ids]
usable = [(sid, rs) for sid, rs in blocks if len(rs) >= 4]
skipped = [(sid, rs) for sid, rs in blocks if len(rs) < 4]
smeta = {s["id"]: s for s in DATA["studies"]}

head(va, 4, ["Study", "Region / respondents", "Risks in common (n)", "Spearman rho", "t", "p (two-tailed)", "Reading", ""],
     [10, 46, 12, 12, 9, 12, 46, 4])
summary_row = {sid: 5 + k for k, (sid, _) in enumerate(usable)}
row = 6 + len(usable) + 2
for sid, rs in usable:
    s = smeta[sid]
    va.cell(row=row, column=1, value=f"{sid}: {s['authors']} {s['year']} ({s['country']})").font = f(bold=True, size=11)
    row += 1
    head(va, row, ["Risk ID", "Risk", "Seed mean RII", "Study RII (direct, averaged)", "Seed rank", "Study rank", "Rank difference"])
    b0 = row + 1
    b1 = b0 + len(rs) - 1
    for i, rid in enumerate(rs, b0):
        put(va, i, 1, rid, bold=True)
        put(va, i, 2, NAMES[rid])
        put(va, i, 3, f"=INDEX({SEED_MEAN},MATCH(A{i},{SEED_IDS},0))", GREEN, fmt="0.0000")
        put(va, i, 4, f'=AVERAGEIFS({rng("E")},{rng("A")},"{sid}",{rng("G")},A{i},{rng("H")},"Direct")', GREEN, fmt="0.000")
        put(va, i, 5, f"=RANK(C{i},$C${b0}:$C${b1},0)+(COUNTIF($C${b0}:$C${b1},C{i})-1)/2", fmt="0.0")
        put(va, i, 6, f"=RANK(D{i},$D${b0}:$D${b1},0)+(COUNTIF($D${b0}:$D${b1},D{i})-1)/2", fmt="0.0")
        put(va, i, 7, f"=E{i}-F{i}", fmt="0.0")
    r_rho, r_n, r_t, r_p = b1 + 1, b1 + 2, b1 + 3, b1 + 4
    for rr, lab, formula, fmt in [
            (r_rho, "Spearman rho (CORREL of the ranks)", f"=CORREL(E{b0}:E{b1},F{b0}:F{b1})", "0.000"),
            (r_n, "n (risks in common)", f"=COUNT(D{b0}:D{b1})", "0"),
            (r_t, "t = rho x SQRT((n-2)/(1-rho^2))", f"=C{r_rho}*SQRT((C{r_n}-2)/(1-C{r_rho}^2))", "0.000"),
            (r_p, "p-value, two-tailed (t-distribution, n-2 df)", f"=TDIST(ABS(C{r_t}),C{r_n}-2,2)", "0.000")]:
        put(va, rr, 2, lab, bold=True)
        put(va, rr, 3, formula, fmt=fmt, bold=True)
    sr = summary_row[sid]
    put(va, sr, 1, sid, bold=True)
    put(va, sr, 2, f"{s['country']}; {s['context']}", wrap=True)
    put(va, sr, 3, f"=C{r_n}", GREEN, fmt="0")
    put(va, sr, 4, f"=C{r_rho}", GREEN, fmt="0.000")
    put(va, sr, 5, f"=C{r_t}", GREEN, fmt="0.000")
    put(va, sr, 6, f"=C{r_p}", GREEN, fmt="0.000")
    put(va, sr, 7, f'=IF(F{sr}<0.05,IF(D{sr}>0,"Agrees (significant at 5%)","Disagrees (significant at 5%)"),'
                   f'IF(D{sr}>=0.5,"Moderate agreement, not significant (small n)",IF(D{sr}>=0.2,"Weak agreement, not significant",'
                   f'IF(D{sr}<=-0.5,"Moderate disagreement, not significant (small n)",IF(D{sr}<=-0.2,"Weak disagreement, not significant","No clear relation")))))',
        wrap=True)
    row = r_p + 3
va.cell(row=row, column=1, value="Held-out studies not tested (fewer than 4 risks in common, too few for a rank correlation):").font = f(bold=True)
for k, (sid, rs) in enumerate(skipped, 1):
    va.cell(row=row + k, column=1, value=f"{sid} ({smeta[sid]['country']}): {len(rs)} risk(s) in common: {', '.join(rs) or 'none'}").font = f(GREY)
row += len(skipped) + 2
for k, t in enumerate([
        "How to read this sheet:",
        "rho = +1: same order as the seed; 0: unrelated; -1: opposite order. The more risks two studies share, the smaller the rho needed for significance (about 0.7 to 0.8 with 7 to 9 risks, about 0.45 with 20).",
        "M10 pools 10 surveys (mostly managers and engineers); M15 asks managers; M14 asks site WORKERS. Who answers changes the ranking, so compare like with like.",
        "This tests the RANKING only. It cannot test probabilities or days of delay; that needs site records (see docs/Site_Data_Collection_Form.docx)."]):
    va.cell(row=row + k, column=1, value=t).font = f(bold=k == 0, color=BLACK if k == 0 else GREY)
va.freeze_panes = "A5"

# ------------------------------------------------------------------ Other_indices
oi = wb.create_sheet("Other_indices")
title(oi, "Values found in the literature but NOT used (different index or group-level only)")
head(oi, 4, ["ID", "Source", "Index", "Values", "Why not used"], [8, 60, 30, 70, 46])
for i, o in enumerate(DATA["other_indices"], 5):
    for c, k in enumerate(["study", "source", "index", "items", "why_not_used"], 1):
        put(oi, i, c, o[k], BLUE if c < 5 else BLACK, wrap=True)


# ------------------------------------------------------------------ Standardised (all studies on one level)
sz = wb.create_sheet("Standardised")
title(sz, "Bringing every study to one level: within-study z-scores",
      "z = (RII - mean RII of the study's factors) / SD of the study's factors. Mean 0 and SD 1 in every study, so scale and respondent differences drop out.")
head(sz, 4, ["Study", "Treatment (judgement)", "Factors (n)", "Mean RII", "SD", "Why"], [9, 20, 11, 11, 10, 70])
TREAT = {"M06": ("Complete", "all 40 factors printed"), "M10": ("Complete", "33 of 35 pooled factors captured"), "M14": ("Complete", "all 23 factors printed"),
         "M15": ("Complete", "all 24 factors printed"), "M01": ("Truncated", "paper lists only the 28 factors with RII > 0.7 (of 38): the lowest are cut off, so its z-scores are distorted; used only in Analysis B")}
ids = [x["id"] for x in DATA["studies"]]
for i, sid in enumerate(ids, 5):
    tr, why = TREAT.get(sid, ("Partial", "only some factors recorded (top-ranked or selected): z-scores would be biased, so not standardised"))
    put(sz, i, 1, sid, bold=True)
    put(sz, i, 2, tr, BLUE, fill=ASSUME_FILL)
    put(sz, i, 3, f'=COUNTIF({rng("A")},A{i})')
    put(sz, i, 4, f'=AVERAGEIF({rng("A")},A{i},{rng("E")})', fmt="0.0000")
    put(sz, i, 5, f'=IF(C{i}>1,SQRT(SUMPRODUCT(({rng("A")}=A{i})*({rng("E")}-D{i})^2)/(C{i}-1)),"")', fmt="0.0000")
    put(sz, i, 6, why, GREY, wrap=True)
t0 = 5 + len(ids) + 2
sz.cell(row=t0, column=1, value="Risk score = mean z of the study's DIRECT factors for the risk. Analysis A = the four complete studies; Analysis B adds M01.").font = f(bold=True)
zs = ["M06", "M10", "M14", "M15", "M01"]
head(sz, t0 + 1, ["Risk", "Name"] + zs + ["Studies (A)", "Pooled z (A)", "Rank (A)", "Studies (B)", "Pooled z (B)", "Rank (B)"],
     None)
for c, w in zip("ABCDEFGHIJKLM", [9, 44, 9, 9, 9, 9, 9, 11, 12, 9, 11, 12, 9]):
    sz.column_dimensions[c].width = max(sz.column_dimensions[c].width or 0, w)
a0, a1 = t0 + 2, t0 + 1 + len(RISKS)
for i, rid in enumerate(RISKS, a0):
    put(sz, i, 1, rid, bold=True)
    put(sz, i, 2, NAMES[rid], wrap=True)
    for j, sid in enumerate(zs):
        col = 3 + j
        put(sz, i, col, f'=IFERROR(AVERAGEIFS({rng("K")},{rng("A")},"{sid}",{rng("G")},$A{i},{rng("H")},"Direct"),"")', GREEN, fmt="0.00")
    put(sz, i, 8, f"=COUNT(C{i}:F{i})")
    put(sz, i, 9, f'=IF(H{i}=0,"",AVERAGE(C{i}:F{i}))', fmt="0.000", bold=True)
    put(sz, i, 10, f'=IF(I{i}="","",RANK(I{i},$I${a0}:$I${a1},0))')
    put(sz, i, 11, f"=COUNT(C{i}:G{i})")
    put(sz, i, 12, f'=IF(K{i}=0,"",AVERAGE(C{i}:G{i}))', fmt="0.000", bold=True)
    put(sz, i, 13, f'=IF(L{i}="","",RANK(L{i},$L${a0}:$L${a1},0))')
nn = a1 + 2
for k, t in enumerate([
        "How to read this sheet:",
        "1. Different rating scales give different RII levels (1-4 studies average about 0.59, 1-5 studies about 0.72 to 0.75). Standardising removes that, so values from different studies can be averaged.",
        "2. Rank agreement WITHIN pairs of studies does not change under standardisation (z keeps the order inside a study); what changes is the pooled value and its uncertainty.",
        "3. Treatment (yellow) is a judgement: only studies whose full factor list is available can be standardised honestly. Change it and the sheet recalculates.",
        "4. Compare Pooled z (A) and (B) with the seed order on Seed_Calculation. They disagree: see docs/Statistical_Analysis_Standardised.docx for the tests (scripts/stats_standardised.py)."]):
    sz.cell(row=nn + k, column=1, value=t).font = f(bold=k == 0, color=BLACK if k == 0 else GREY)
sz.freeze_panes = "A5"

# ------------------------------------------------------------------ How it works (first sheet)
hw = wb.create_sheet("How it works", 0)
hw.column_dimensions["A"].width = 4
hw.column_dimensions["B"].width = 120
title(hw, "How Bricks-ly uses RII values from published research")
lines = [
    ("h", "What an RII is"),
    ("t", "Relative Importance Index = sum of all ratings / (highest possible rating x number of respondents). Survey respondents rate each delay factor (e.g. 1-5); the RII runs from about 0.2 to 1. A higher RII means respondents judged the factor more important."),
    ("t", "An RII is NOT a probability and NOT a number of days of delay. It can only say which risks people consider more important than others."),
    ("h", "Where the values come from"),
    ("t", "Sheet 'Studies' lists 11 published surveys; sheet 'RII_Data' holds every value recorded from them, as printed, with the table and page. Every factor is mapped to a library risk (only 4 unclear or group-level factors are left unmapped)."),
    ("t", "Seed studies (used by the model): RII surveys from India and Sri Lanka, closest to Indian masonry sites: M01 (India, masonry), A02 (India, masonry), M06 (India), A14 (Sri Lanka), A15 (Sri Lanka)."),
    ("t", "Validation studies (held out, never used by the model): M10 (Middle East meta-analysis of 10 surveys), M15 (Brazil), M14 (Indonesia, workers), plus A10, A21 and HAS25 with too few overlapping risks to test."),
    ("h", "Step by step (sheet 'Seed_Calculation' does this with live formulas)"),
    ("t", "1. Map each factor to one of the tool's library risks (the risk library; only the factors entered in the seed dataset are mapped, not every factor in each paper). Direct = names the same cause (e.g. 'Material shortages' -> R-MAT). Related = overlaps but is broader, narrower or an outcome (e.g. 'Accidents' -> R-SAFE); Related values are shown but not used."),
    ("t", "2. For each risk and seed study, average that study's Direct values, so one study counts once even if it lists two matching factors."),
    ("t", "3. Seed mean RII for the risk = average across the seed studies that have a value."),
    ("t", "4. Rank the seeded risks by seed mean RII and split them into five equal-count bands (with 30 seeded risks, six per band): top band -> class 5, ... bottom band -> class 1. Risks with no seed-study value (for example R-TRN, R-PRD, R-GOV) are not seeded."),
    ("t", "5. In the app, a library risk with no probability or delay entered yet is placed on the 5x5 matrix at (class, class), drawn with a dashed outline and labelled 'literature seed'."),
    ("t", "6. If a site-fact rule flags the risk (e.g. monsoon overlaps the work), its probability class goes up by one, to a maximum of 5."),
    ("t", "7. As soon as the user enters a probability and a delay range (with a Source), the seed is replaced by those numbers."),
    ("h", "Worked example: R-WX (rain / monsoon stoppage)"),
    ("f", f'="Study values: M01 "&TEXT(Seed_Calculation!C{first + 6},"0.000")&", A02 "&TEXT(Seed_Calculation!D{first + 6},"0.000")&", M06 "&TEXT(Seed_Calculation!E{first + 6},"0.000")&", A14 "&TEXT(Seed_Calculation!F{first + 6},"0.000")'),
    ("f", f'="Seed mean RII = "&TEXT(Seed_Calculation!{MEAN_COL}{first + 6},"0.0000")&" -> rank "&Seed_Calculation!{L(cN + 3)}{first + 6}&" of "&Seed_Calculation!C{nrow}&" -> class "&Seed_Calculation!{L(cN + 4)}{first + 6}&" -> matrix cell "&Seed_Calculation!{L(cN + 5)}{first + 6}'),
    ("h", "What the validation shows (sheet 'Validation')"),
    ("t", "Spearman rank correlation between the seed order and each held-out study, over the risks both cover. Read the 'Reading' column; with only 7-9 shared risks, a result must be strong to be statistically significant."),
    ("h", "Limits to state in the report"),
    ("t", "RII is a perception of importance, not a measured frequency or delay. Studies differ in country, project type, respondents (managers vs workers) and rating scale (1-4 vs 1-5). Several values come from research notes rather than a re-read of the PDF (column 'How the values were obtained')."),
    ("t", "Most seed mean RIIs lie within a few hundredths of each other, so the order of the middle risks is fragile. The seed is a starting point until expert-survey or site data replace it."),
    ("h", "Colour key"),
    ("t", "Blue text = value copied from a paper. Black = formula. Green = formula that pulls from another sheet. Yellow fill = a judgement you may change (study role, factor mapping); the sheets recalculate."),
    ("t", "Built by scripts/build_rii_workbook.py from data/literature_seed.json, the same file the app reads. Edit the JSON and rebuild to keep the app and this workbook in step."),
]
r = 3
for kind, text in lines:
    c = hw.cell(row=r, column=1 if kind == "h" else 2, value=text)
    if kind == "h":
        r += 1
        c = hw.cell(row=r, column=1, value=text)
        hw.cell(row=r - 1, column=1).value = None
        c.font = f(bold=True, size=11)
    else:
        c.font = f(GREEN if kind == "f" else BLACK)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title in ("RII_Data",)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(f"wrote {OUT} | studies {len(DATA['studies'])} | rows {len(rows)} | validation blocks {[s for s, _ in usable]} | skipped {[s for s, _ in skipped]}")
