"""Build the ranking Excel from data/openalex_corpus.jsonl.

Ranking is a transparent, formula-based COMPOSITE of citation count, recency, open-access status and
whether the PDF was actually downloaded -- never a fabricated "quality" judgment. The exact formula is
written into the Methodology sheet of the workbook itself, next to the numbers it produced, so nothing
here is a mystery score. cited_by_count and is_oa are OpenAlex's own reported values, not derived.
"""
from __future__ import annotations

import json
import math
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.cell import cell as _cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# OpenAlex abstracts occasionally contain control characters that Excel cannot store; drop them rather than fail the export
_ILLEGAL = __import__("re").compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_orig_check = _cell.Cell.check_string
_cell.Cell.check_string = lambda self, value: _orig_check(self, _ILLEGAL.sub("", value) if isinstance(value, str) else value)

ROOT = Path(__file__).resolve().parents[1]
JSONL = ROOT / "data" / "openalex_corpus.jsonl"
OUT = ROOT / "Research_Papers_Ranking.xlsx"

W_CITE, W_RECENCY, W_OA, W_DL = 0.45, 0.20, 0.15, 0.20
CURRENT_YEAR = datetime.now().year
MIN_YEAR = 1995


def load_records() -> list[dict]:
    recs = []
    with open(JSONL, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                recs.append(json.loads(line))
    return recs


def obtained(r: dict) -> bool:
    """A verified PDF was fetched for this paper (it may since have been moved to the archive outside the project)."""
    return bool(r.get("downloaded") or r.get("archived_pdf"))


def score(recs: list[dict]) -> list[dict]:
    max_log_cite = max((math.log1p(r.get("cited_by_count") or 0) for r in recs), default=1) or 1
    for r in recs:
        cite = math.log1p(r.get("cited_by_count") or 0) / max_log_cite
        yr = r.get("year") or MIN_YEAR
        recency = max(0.0, min(1.0, (yr - MIN_YEAR) / max(1, (CURRENT_YEAR - MIN_YEAR))))
        oa = 1.0 if r.get("is_oa") else 0.0
        dl = 1.0 if obtained(r) else 0.0
        r["_citation_component"] = round(cite, 4)
        r["_recency_component"] = round(recency, 4)
        r["_oa_component"] = oa
        r["_downloaded_component"] = dl
        r["_composite_score"] = round(W_CITE * cite + W_RECENCY * recency + W_OA * oa + W_DL * dl, 4)
    recs.sort(key=lambda r: -r["_composite_score"])
    for i, r in enumerate(recs, 1):
        r["_rank"] = i
    return recs


def autosize(ws, widths: dict[int, int]):
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w


def build(recs: list[dict]):
    wb = openpyxl.Workbook()

    # ---- Methodology sheet first, so the ranking is never a black box ----
    ws0 = wb.active
    ws0.title = "Methodology"
    lines = [
        ["Research paper ranking: methodology", ""],
        ["", ""],
        ["Source", "OpenAlex (https://openalex.org), a free, open scholarly index. Every row's 'OpenAlex ID' and DOI can be independently verified there."],
        ["What is real vs. derived", "Title, authors, year, venue, DOI, citation count and open-access status are OpenAlex's own reported values, copied as-is. "
         "'PDF obtained' is Yes only when this project's harvester actually fetched a file that started with the PDF signature (%PDF) and saved it -- see the Local path column. "
         "The harvested PDFs were later moved to an archive folder outside the project (FYP_backups/literature_archive), so Local path points there; the curated papers in Research_Papers/1-5 were not moved. "
         "Everything under '_component' or 'Composite score' below is DERIVED by the formula stated here, not reported by OpenAlex."],
        ["Composite score formula", f"{W_CITE} x citation_component + {W_RECENCY} x recency_component + {W_OA} x oa_component + {W_DL} x downloaded_component"],
        ["citation_component", "log(1 + citations) / log(1 + max citations in this corpus) -- a log scale so a handful of very highly cited papers do not dominate the ranking."],
        ["recency_component", f"(year - {MIN_YEAR}) / ({CURRENT_YEAR} - {MIN_YEAR}), clipped to [0,1]. Newer papers score higher; this is a recency preference, not a quality judgment."],
        ["oa_component", "1 if OpenAlex reports the work as open access, else 0. Reflects accessibility for this project, not quality."],
        ["downloaded_component", "1 if a verified PDF was obtained for the paper (now kept in the archive folder), else 0. Reflects availability of the full text to this project, not the paper's merit."],
        ["What this ranking is NOT", "It is not a peer-review-quality score, an impact-factor proxy, or a claim about correctness. Two 0-citation papers with different years/OA status can rank apart purely on that basis. "
         "Always read a paper before relying on a specific number from it; the rank only helps decide reading order."],
        ["Corpus scope", "Records kept after screening (scripts/build_literature_db.py, rules in docs/Literature_Database_Report.md): construction works about delay causes, "
         "risk identification/assessment (RII, probability x impact, matrix), labour productivity and masonry, and rule-based/AI risk tools. Monte Carlo, EMV, risk response, "
         "cost-overrun and decision-support topics were dropped. 'Topic bucket' is the harvest search that first found a record (provenance only); the 'By bucket' sheet lists the queries."],
        ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")],
        ["Total records", len(recs)],
        ["PDF obtained and verified", sum(1 for r in recs if obtained(r))],
        ["Open access per OpenAlex", sum(1 for r in recs if r.get("is_oa"))],
        ["No PDF obtained (metadata / abstract only)", sum(1 for r in recs if not obtained(r))],
    ]
    for row in lines:
        ws0.append(row)
    ws0.column_dimensions["A"].width = 34
    ws0.column_dimensions["B"].width = 110
    for row in ws0.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws0["A1"].font = Font(bold=True, size=14)

    # ---- Main ranking sheet ----
    ws = wb.create_sheet("Ranked Papers")
    head = ["Rank", "Composite score", "Title", "Authors", "Year", "Venue", "DOI", "Citations",
            "Open access", "PDF obtained", "Local path", "Topic bucket", "Matched query",
            "OpenAlex ID", "Landing page URL", "Citation component", "Recency component", "OA component",
            "Downloaded component", "Abstract"]
    ws.append(head)
    for r in recs:
        ws.append([
            r["_rank"], r["_composite_score"], r["title"], "; ".join(r.get("authors") or []), r.get("year"),
            r.get("venue"), r.get("doi"), r.get("cited_by_count"), "Yes" if r.get("is_oa") else "No",
            "Yes" if obtained(r) else "No", r.get("local_path") or r.get("archived_path") or "", r.get("topic_bucket"),
            r.get("matched_query"), r.get("id"), r.get("landing_url"),
            r["_citation_component"], r["_recency_component"], r["_oa_component"], r["_downloaded_component"],
            (r.get("abstract") or "")[:400],
        ])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    widths = {1: 6, 2: 10, 3: 46, 4: 30, 5: 6, 6: 26, 7: 22, 8: 9, 9: 9, 10: 12, 11: 30, 12: 30, 13: 26,
             14: 26, 15: 30, 16: 10, 17: 10, 18: 8, 19: 10, 20: 50}
    autosize(ws, widths)
    ws.freeze_panes = "D2"
    tbl = Table(displayName="RankedPapers", ref=f"A1:{get_column_letter(len(head))}{len(recs) + 1}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)

    # ---- By-bucket summary ----
    ws2 = wb.create_sheet("By bucket")
    ws2.append(["Topic bucket", "Search queries used", "Papers found", "PDF obtained", "Open access", "Median citations"])
    from collections import defaultdict

    by_bucket: dict[str, list[dict]] = defaultdict(list)
    for r in recs:
        by_bucket[r["topic_bucket"]].append(r)
    for bucket, items in sorted(by_bucket.items()):
        qs = sorted({r["matched_query"] for r in items})
        cites = sorted(r.get("cited_by_count") or 0 for r in items)
        med = cites[len(cites) // 2] if cites else 0
        ws2.append([bucket, "; ".join(qs), len(items), sum(1 for r in items if obtained(r)),
                   sum(1 for r in items if r["is_oa"]), med])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    autosize(ws2, {1: 34, 2: 70, 3: 12, 4: 12, 5: 12, 6: 14})
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    target = OUT
    try:
        wb.save(target)
    except PermissionError:                      # open in Excel: save beside it instead of failing
        target = OUT.with_name(OUT.stem + "_new.xlsx")
        wb.save(target)
    return target


if __name__ == "__main__":
    recs = score(load_records())
    path = build(recs)
    print(f"Wrote {len(recs)} ranked records to {path}")
