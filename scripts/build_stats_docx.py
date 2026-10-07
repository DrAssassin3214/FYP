"""Combine docs/Literature_Database_Report.md and docs/Corpus_Statistical_Analysis.md into one Word file.

Handles only what those two generated files use: '#'/'##' headings, paragraphs, '* ' bullets, pipe tables, **bold**, `code`.
Run: python scripts/build_stats_docx.py   (after build_literature_db.py, corpus_statistics.py, stats_corpus.py)
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Literature_Database_and_Statistical_Analysis.docx"
PARTS = [("Part A. Literature database: screening and size", ROOT / "docs" / "Literature_Database_Report.md"),
         ("Part B. Statistical analysis of the database", ROOT / "docs" / "Corpus_Statistical_Analysis.md")]


def runs(p, text: str) -> None:
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if tok.startswith("**"):
            p.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`"):
            r = p.add_run(tok[1:-1])
            r.font.name = "Consolas"
        elif tok:
            p.add_run(tok)


def render(doc: Document, md: str) -> None:
    lines, i = md.splitlines(), 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# "):
            doc.add_heading(ln[2:], 2)
        elif ln.startswith("## "):
            doc.add_heading(ln[3:], 3)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[-:| ]+\|$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, val in enumerate(row[:len(rows[0])]):
                    cell = t.cell(r, c)
                    cell.text = ""
                    runs(cell.paragraphs[0], val)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(8.5)
                        run.bold = run.bold or r == 0
                    if c > 0 and re.fullmatch(r"[-+<>0-9.,%() ]+", val):
                        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            doc.add_paragraph()
            continue
        elif ln.startswith("* "):
            runs(doc.add_paragraph(style="List Bullet"), ln[2:])
        elif ln.strip():
            runs(doc.add_paragraph(), ln)
        i += 1


def main() -> int:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading("Literature Database and Statistical Analysis", 0)
    doc.add_paragraph("Bricks-ly FYP. All figures are computed by scripts in the project from data/literature.db (OpenAlex metadata, "
                      "screening rules 2026-10-07.1). Themes, tiers and method/region shares are derived by keyword rules, not reported by OpenAlex. "
                      "The relevance (precision) of the screening has not yet been measured: it needs the judged audit sample.")
    for title, path in PARTS:
        doc.add_heading(title, 1)
        render(doc, path.read_text(encoding="utf-8"))
    out = OUT
    try:
        doc.save(out)
    except PermissionError:                      # the file is open in Word: save beside it instead of failing
        out = OUT.with_name(OUT.stem + "_new.docx")
        doc.save(out)
    print("wrote", out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
