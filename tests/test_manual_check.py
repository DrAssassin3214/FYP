"""Manual-check workbook: structure, the pasted tool values, failure on a tampered value, the API route and the CLI."""
import copy
import io
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import cli, service
from app.reporting.manual_check import manual_check_xlsx

SHEETS = ["README", "Inputs", "Risks", "Totals", "Matrix", "Responses"]


def _wb(case, **kw):
    return load_workbook(io.BytesIO(manual_check_xlsx(case, **kw)))


def test_sheets_formulas_and_pasted_tool_values():
    case = service.example_analysis_case()
    res = service.run_analysis(case)
    wb = _wb(case, result=res)
    assert wb.sheetnames == SHEETS
    n_f = sum(1 for s in wb for row in s.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    assert n_f > 150
    r = wb["Risks"]
    for i, e in enumerate(res["event_emv"]):
        assert r[f"H{4+i}"].value == pytest.approx(e["emv"]) and r[f"G{4+i}"].value == pytest.approx(e["expected_delay_if_occurs_days"])
    assert wb["Totals"]["C4"].value == pytest.approx(res["summary"]["expected_delay"])
    p = wb["Responses"]
    for i, c in enumerate(res["mitigation_checks"]):
        assert p[f"P{4+i}"].value == pytest.approx(c["net_benefit"])


def test_correlated_case_marks_independence_row_na():
    case = service.example_analysis_case()
    ids = [x["id"] for x in case["risks"]]
    case["simulation"]["correlation"] = [{"a": ids[0], "b": ids[1], "rho": 0.6}]
    assert _wb(case)["Totals"]["F5"].value == "n/a"
    assert _wb(service.example_analysis_case())["Totals"]["F5"].value.startswith("=IF(")


def test_case_without_responses_builds():
    case = service.example_analysis_case()
    case["mitigations"], case["options"] = [], []
    assert "Responses" in _wb(case).sheetnames


def test_invalid_case_raises_case_error():
    bad = service.example_analysis_case()
    del bad["cost"]
    with pytest.raises(service.CaseError):
        manual_check_xlsx(bad)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_recalculated_all_pass_and_tamper_fails(tmp_path):
    def recalc(case, res=None):
        p = tmp_path / "m.xlsx"
        p.write_bytes(manual_check_xlsx(case, result=res))
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "o"), str(p)],
                       check=True, capture_output=True, timeout=120)
        return load_workbook(tmp_path / "o" / "m.xlsx", data_only=True)

    case = service.example_analysis_case()
    res = service.run_analysis(case)
    assert recalc(case, res)["README"]["B12"].value.startswith("ALL ")
    bad = copy.deepcopy(res)
    bad["event_emv"][0]["emv"] += 1.0
    assert "FAIL" in recalc(case, bad)["README"]["B12"].value


def test_api_route():
    from gui.api import create_app

    c = create_app().test_client()
    ok = c.post("/api/manual-check-xlsx", json=service.example_analysis_case())
    assert ok.status_code == 200 and ok.data[:2] == b"PK" and "spreadsheetml" in ok.mimetype
    assert "manual_check.xlsx" in ok.headers["Content-Disposition"]
    bad = service.example_analysis_case()
    del bad["cost"]
    assert c.post("/api/manual-check-xlsx", json=bad).status_code == 422


def test_cli(tmp_path, capsys):
    case = tmp_path / "c.json"
    case.write_text(__import__("json").dumps(service.example_analysis_case()))
    out = tmp_path / "x.xlsx"
    assert cli.main(["manual-check", str(case), "--out", str(out)]) == 0 and out.read_bytes()[:2] == b"PK"
    assert cli.main(["manual-check", str(tmp_path / "missing.json"), "--out", str(out)]) == 2
    assert cli.main(["manual-check", str(case), "--out", str(tmp_path / "nodir" / "x.xlsx")]) == 2
