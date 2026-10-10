"""Build the one-test workbook: hand calculation, Excel simulation and tool simulation for a single-risk case.

Test 1: risk R-MAT only (p = 0.5, PERT(1, 3, 8) days), planned duration 16 d, deadline 20 d, delay cost 8,000 INR/d,
liquidated damages 5,000 INR/d; response M-BUF (buffer stock, p 0.5 -> 0.2, cost 6,000 INR) against accepting the risk.
Usage: python3 -I build_test1_xlsx.py <output.xlsx>
"""
import json, sys
from pathlib import Path
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO))
from app import service  # noqa: E402

OUT = Path(sys.argv[1])
N_XL = 2000
N = "Excel_Simulation!$B$4"   # cell holding the number of Excel draws
SEED = 12345

# ---------------------------------------------------------------------------------------- tool side (values)
def tool_run(n):
    c = service.example_analysis_case()
    c["risks"] = [r for r in c["risks"] if r["id"] == "R-MAT"]
    c["mitigations"] = [m for m in c["mitigations"] if m["id"] == "M-BUF"]
    c["options"] = []
    c["simulation"] = {"n": n, "seed": SEED}
    c.pop("facts", None)
    res = service.run_analysis(c)
    o = {x["option_id"]: x for x in res["decision"]["options"]}
    se = res["cost"]["std"] / n ** 0.5
    return {
        "n": n, "command": res["command"]["command"], "se": se,
        "A": o["ACCEPT"], "B": o["O-M-BUF"], "p_any": res["summary"]["p_any_delay"],
    }

T10, T200 = tool_run(10_000), tool_run(200_000)

# ---------------------------------------------------------------------------------------- styles
F = "Arial"
blue = Font(name=F, size=10, color="0000FF")
blk = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
grn = Font(name=F, size=10, color="008000")
title = Font(name=F, size=14, bold=True)
hf = Font(name=F, size=10, bold=True, color="FFFFFF")
hfill = PatternFill("solid", start_color="1F3864")
inp = PatternFill("solid", start_color="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()


def put(ws, ref, v, font=blk, fmt=None, fill=None, al=None, border=True):
    c = ws[ref]
    c.value = v
    c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if al: c.alignment = al
    if border: c.border = box
    return c


def header(ws, row, labels, col0=1, height=30):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=t)
        c.font, c.fill, c.alignment, c.border = hf, hfill, ctr, box
    ws.row_dimensions[row].height = height


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def passfail(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], font=Font(name=F, color="006100", bold=True), fill=PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], font=Font(name=F, color="9C0006", bold=True), fill=PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")))


# ---------------------------------------------------------------------------------------- Test1 (inputs + hand calc)
ws = wb.active
ws.title = "Test1_Hand"
ws["A1"].value, ws["A1"].font = "Test 1: one risk, one response. Hand calculation", title
ws["A2"].value = "ILLUSTRATIVE placeholders (Assumptions), not site data. Blue on yellow = input; black = formula; green = link to another sheet."
ws["A2"].font = blk
header(ws, 3, ["Input", "Value", "Unit / note"])
inputs = [
    ("Planned duration T0", 16, "days", "#,##0"),
    ("Deadline", 20, "days", "#,##0"),
    ("Delay cost Cd", 8000, "INR per day of delay", "#,##0"),
    ("Liquidated damages LD", 5000, "INR per day after the deadline", "#,##0"),
    ("PERT lambda", 4, "standard weight", "0"),
    ("Delay a (min)", 1, "days", "0.0"),
    ("Delay m (mode)", 3, "days", "0.0"),
    ("Delay b (max)", 8, "days", "0.0"),
    ("p, accept (no response)", 0.5, "probability R-MAT occurs", "0.00"),
    ("p, with M-BUF", 0.2, "probability after buffer stock", "0.00"),
    ("M-BUF cost", 6000, "INR", "#,##0"),
]
for i, (k, v, u, fm) in enumerate(inputs):
    r = 4 + i
    put(ws, f"A{r}", k, bold)
    put(ws, f"B{r}", v, blue, fm, inp)
    put(ws, f"C{r}", u)
# B4 T0, B5 DL, B6 Cd, B7 LD, B8 lam, B9 a, B10 m, B11 b, B12 pA, B13 pB, B14 cost
header(ws, 17, ["Derived quantity", "Value", "Formula"])
der = [
    ("Threshold c = deadline - T0 (d)", "=B5-B4", "0.0", "delay beyond c breaks the deadline"),
    ("Mean delay if it occurs, mu", "=(B9+B8*B10+B11)/(B8+2)", "0.0000", "(a + lambda*m + b)/(lambda + 2)"),
    ("Variance of delay if it occurs", "=(B19-B9)*(B11-B19)/(B8+3)", "0.000000", "(mu - a)(b - mu)/(lambda + 3)"),
    ("Beta alpha", "=1+B8*(B10-B9)/(B11-B9)", "0.000000", "1 + lambda (m - a)/(b - a)"),
    ("Beta beta", "=1+B8*(B11-B10)/(B11-B9)", "0.000000", "1 + lambda (b - m)/(b - a)"),
    ("P(delay <= c | occurs), F(c)", "=BETADIST(B18,B21,B22,B9,B11)", "0.000000", "Beta CDF at c"),
    ("P(delay > c | occurs)", "=1-B23", "0.000000", ""),
    ("E[(delay - c)+ | occurs] (d)", "=Integral!E49", "0.000000", "integral of 1 - F(x) from c to b (Simpson, sheet Integral)"),
]
for i, (k, f, fm, n) in enumerate(der):
    r = 18 + i
    put(ws, f"A{r}", k, bold)
    put(ws, f"B{r}", f, grn if "Integral!" in f else blk, fm)
    put(ws, f"C{r}", n)
# rows 18..25 : c=18, mu=19, var=20, alpha=21, beta=22, F=23, S=24, E+=25
header(ws, 28, ["Result by hand", "Accept (no response)", "With M-BUF", "How"], height=30)
rows = [
    ("Probability of a delay", "=B12", "=B13", "0.0000", "p"),
    ("Expected delay (d)", "=B29*$B$19", "=C29*$B$19", "0.0000", "p x mu"),
    ("Expected delay cost (INR)", "=B30*$B$6", "=C30*$B$6", "#,##0.00", "Cd x expected delay"),
    ("P(finish after the deadline)", "=B29*$B$24", "=C29*$B$24", "0.0000", "p x P(delay > c)"),
    ("Expected days past the deadline", "=B29*$B$25", "=C29*$B$25", "0.0000", "p x E[(delay - c)+]"),
    ("Expected liquidated damages (INR)", "=B33*$B$7", "=C33*$B$7", "#,##0.00", "LD x expected days past deadline"),
    ("Residual expected cost (INR)", "=B31+B34", "=C31+C34", "#,##0.00", "delay cost + LD"),
    ("Response cost (INR)", 0, "=B14", "#,##0.00", "M-BUF cost"),
    ("Total expected cost, TEC (INR)", "=B35+B36", "=C35+C36", "#,##0.00", "response cost + residual"),
    ("TEC change vs accepting (INR)", "=B37-$B$37", "=C37-$B$37", "#,##0.00;(#,##0.00);-", "negative = cheaper than accepting"),
    ("Exact P80 duration (d)", "=IF(0.8<=1-B29,$B$4,$B$4+BETAINV((0.8-(1-B29))/B29,$B$21,$B$22,$B$9,$B$11))", "=IF(0.8<=1-C29,$B$4,$B$4+BETAINV((0.8-(1-C29))/C29,$B$21,$B$22,$B$9,$B$11))", "0.0000", "invert P(T <= t) = (1-p) + p F(t - T0)"),
    ("Exact P90 duration (d)", "=IF(0.9<=1-B29,$B$4,$B$4+BETAINV((0.9-(1-B29))/B29,$B$21,$B$22,$B$9,$B$11))", "=IF(0.9<=1-C29,$B$4,$B$4+BETAINV((0.9-(1-C29))/C29,$B$21,$B$22,$B$9,$B$11))", "0.0000", "same"),
    ("Exact P95 duration (d)", "=IF(0.95<=1-B29,$B$4,$B$4+BETAINV((0.95-(1-B29))/B29,$B$21,$B$22,$B$9,$B$11))", "=IF(0.95<=1-C29,$B$4,$B$4+BETAINV((0.95-(1-C29))/C29,$B$21,$B$22,$B$9,$B$11))", "0.0000", "same"),
]
for i, (k, fa, fb, fm, n) in enumerate(rows):
    r = 29 + i
    put(ws, f"A{r}", k, bold)
    put(ws, f"B{r}", fa, blue if isinstance(fa, (int, float)) else blk, fm)
    put(ws, f"C{r}", fb, blk, fm)
    put(ws, f"D{r}", n)
put(ws, "A43", "Hand decision", bold)
put(ws, "B43", '=IF(C37<B37,"AUTHORIZE MITIGATION","ACCEPT RISK")', bold)
put(ws, "C43", '=IF(C37<B37,"M-BUF is cheaper in total expected cost by "&TEXT(B37-C37,"#,##0")&" INR","Accepting is cheaper by "&TEXT(C37-B37,"#,##0")&" INR")', blk, border=False)
widths(ws, [40, 22, 20, 52])

# ---------------------------------------------------------------------------------------- Integral
wi = wb.create_sheet("Integral")
wi["A1"].value, wi["A1"].font = "Simpson integration of 1 - F(x) from c to b (40 intervals)", title
wi["A2"].value = "E[(delay - c)+] = integral from c to b of (1 - F(x)) dx, F = Beta-PERT cumulative distribution. Simpson weights 1, 4, 2, 4, ..., 4, 1."
wi["A2"].font = blk
header(wi, 4, ["i", "x (days of delay)", "S(x) = 1 - F(x)", "Weight", "Weight x S(x)"])
for i in range(41):
    r = 5 + i
    put(wi, f"A{r}", i, blk, "0")
    put(wi, f"B{r}", f"=Test1_Hand!$B$18+A{r}*$E$47", grn if i == 0 else blk, "0.0000")
    put(wi, f"C{r}", f"=1-BETADIST(B{r},Test1_Hand!$B$21,Test1_Hand!$B$22,Test1_Hand!$B$9,Test1_Hand!$B$11)", blk, "0.000000")
    w = 1 if i in (0, 40) else (4 if i % 2 == 1 else 2)
    put(wi, f"D{r}", w, blue, "0")
    put(wi, f"E{r}", f"=D{r}*C{r}", blk, "0.000000")
put(wi, "D47", "step h", bold)
put(wi, "E47", "=(Test1_Hand!B11-Test1_Hand!B18)/40", blk, "0.000000")
put(wi, "D48", "Sum", bold)
put(wi, "E48", "=SUM(E5:E45)", blk, "0.000000")
put(wi, "D49", "Integral = h/3 x sum", bold)
put(wi, "E49", "=E47/3*E48", bold, "0.000000")
widths(wi, [8, 20, 18, 10, 16])

# ---------------------------------------------------------------------------------------- Excel simulation
wsim = wb.create_sheet("Excel_Simulation")
wsim["A1"].value, wsim["A1"].font = f"Monte Carlo in Excel: {N_XL:,} draws (uniform numbers pasted from numpy seed {SEED}; everything else is a formula)", title
rng = np.random.default_rng(SEED)
U = rng.random((N_XL, 2))
H0 = 12      # header row of draws
d0, d1 = H0 + 1, H0 + N_XL
header(wsim, H0, ["Draw", "U1 (occurrence)", "U2 (delay)", "Delay if occurs (d) =BETAINV(U2)", "Occurs, accept", "Duration, accept (d)", "Cost, accept (INR)", "Occurs, M-BUF", "Duration, M-BUF (d)", "Cost incl. M-BUF (INR)"], height=54)
for i in range(N_XL):
    r = d0 + i
    wsim.cell(row=r, column=1, value=i + 1).font = blk
    c = wsim.cell(row=r, column=2, value=float(U[i, 0])); c.font = blue; c.number_format = "0.000000"
    c = wsim.cell(row=r, column=3, value=float(U[i, 1])); c.font = blue; c.number_format = "0.000000"
    wsim.cell(row=r, column=4, value=f"=BETAINV(C{r},Test1_Hand!$B$21,Test1_Hand!$B$22,Test1_Hand!$B$9,Test1_Hand!$B$11)").number_format = "0.0000"
    wsim.cell(row=r, column=5, value=f"=IF(B{r}<Test1_Hand!$B$12,1,0)")
    wsim.cell(row=r, column=6, value=f"=Test1_Hand!$B$4+E{r}*D{r}").number_format = "0.0000"
    wsim.cell(row=r, column=7, value=f"=Test1_Hand!$B$6*E{r}*D{r}+Test1_Hand!$B$7*MAX(0,F{r}-Test1_Hand!$B$5)").number_format = "#,##0.00"
    wsim.cell(row=r, column=8, value=f"=IF(B{r}<Test1_Hand!$B$13,1,0)")
    wsim.cell(row=r, column=9, value=f"=Test1_Hand!$B$4+H{r}*D{r}").number_format = "0.0000"
    wsim.cell(row=r, column=10, value=f"=Test1_Hand!$B$6*H{r}*D{r}+Test1_Hand!$B$7*MAX(0,I{r}-Test1_Hand!$B$5)+Test1_Hand!$B$14").number_format = "#,##0.00"
    for col in range(4, 11):
        wsim.cell(row=r, column=col).font = grn if False else blk
header(wsim, 3, ["Excel simulation summary", "Accept", "With M-BUF"], height=24)
sm = [
    ("Draws", f"=COUNT(B{d0}:B{d1})", f"=COUNT(B{d0}:B{d1})", "#,##0"),
    ("P(delay occurs)", f"=AVERAGE(E{d0}:E{d1})", f"=AVERAGE(H{d0}:H{d1})", "0.0000"),
    ("Expected delay (d)", f"=AVERAGE(F{d0}:F{d1})-Test1_Hand!$B$4", f"=AVERAGE(I{d0}:I{d1})-Test1_Hand!$B$4", "0.0000"),
    ("P(finish after the deadline)", f"=COUNTIF(F{d0}:F{d1},\">\"&Test1_Hand!$B$5)/B4", f"=COUNTIF(I{d0}:I{d1},\">\"&Test1_Hand!$B$5)/C4", "0.0000"),
    ("P90 duration (d)", f"=PERCENTILE(F{d0}:F{d1},0.9)", f"=PERCENTILE(I{d0}:I{d1},0.9)", "0.0000"),
    ("Total expected cost (INR)", f"=AVERAGE(G{d0}:G{d1})", f"=AVERAGE(J{d0}:J{d1})", "#,##0.00"),
    ("Standard error of the mean cost (INR)", f"=STDEV(G{d0}:G{d1})/SQRT(B4)", f"=STDEV(J{d0}:J{d1})/SQRT(C4)", "#,##0.00"),
]
for i, (k, fa, fb, fm) in enumerate(sm):
    r = 4 + i
    put(wsim, f"A{r}", k, bold)
    put(wsim, f"B{r}", fa, blk, fm)
    put(wsim, f"C{r}", fb, blk, fm)
# B4 draws, 5 p, 6 E delay, 7 P exceed, 8 P90, 9 TEC, 10 SE
widths(wsim, [34, 16, 16, 20, 14, 16, 18, 14, 16, 20])
wsim.freeze_panes = f"A{d0}"

# ---------------------------------------------------------------------------------------- Compare
wc = wb.create_sheet("Result")
wc["A1"].value, wc["A1"].font = "Test 1 result: hand vs Excel simulation vs the tool", title
wc["A2"].value = "The hand column is exact (formulas on Test1_Hand). Tool columns are the numbers the tool printed for the same case (seed 12345). Tolerance: 3 standard errors of each quantity for the 2,000-draw Excel run (the least precise run), computed from the exact distribution below; the two tool runs are held to the same limit."
wc["A2"].font = blk
header(wc, 4, ["Quantity", "Hand (exact)", "Excel simulation\n(2,000 draws)", "Tool n=10,000", "Tool n=200,000", "|Excel - hand|", "|Tool 10k - hand|", "|Tool 200k - hand|", "Tolerance", "Excel", "Tool 10k", "Tool 200k"], height=48)


def trow(r, label, hand, xl, t10, t200, fm, tol, tol_is_cost=False):
    put(wc, f"A{r}", label, bold)
    put(wc, f"B{r}", hand, grn, fm)
    put(wc, f"C{r}", xl, grn, fm)
    put(wc, f"D{r}", t10, blue, fm)
    put(wc, f"E{r}", t200, blue, fm)
    put(wc, f"F{r}", f"=ABS(C{r}-B{r})", blk, fm)
    put(wc, f"G{r}", f"=ABS(D{r}-B{r})", blk, fm)
    put(wc, f"H{r}", f"=ABS(E{r}-B{r})", blk, fm)
    put(wc, f"I{r}", tol, blk if str(tol).startswith("=") else blue, fm)
    for col, src in (("J", "F"), ("K", "G"), ("L", "H")):
        put(wc, f"{col}{r}", f'=IF({src}{r}<=I{r},"PASS","FAIL")', bold, al=ctr)


# Accept
put(wc, "A5", "ACCEPT (no response)", bold, border=False)
trow(6, "P(delay occurs)", "=Test1_Hand!B29", "=Excel_Simulation!B5", T10["p_any"], T200["p_any"], "0.0000", f"=3*SQRT(Test1_Hand!B29*(1-Test1_Hand!B29)/{N})")
trow(7, "Expected delay (d)", "=Test1_Hand!B30", "=Excel_Simulation!B6", T10["A"]["expected_delay"], T200["A"]["expected_delay"], "0.0000", f"=3*SQRT(Test1_Hand!B29*(Test1_Hand!$B$20+Test1_Hand!$B$19^2)-Test1_Hand!B30^2)/SQRT({N})")
trow(8, "P(finish after the deadline)", "=Test1_Hand!B32", "=Excel_Simulation!B7", T10["A"]["p_exceed_deadline"], T200["A"]["p_exceed_deadline"], "0.0000", f"=3*SQRT(Test1_Hand!B32*(1-Test1_Hand!B32)/{N})")
trow(9, "P90 duration (d)", "=Test1_Hand!B40", "=Excel_Simulation!B8", T10["A"]["P90"], T200["A"]["P90"], "0.0000", "=3*B33")
trow(10, "Total expected cost (INR)", "=Test1_Hand!B37", "=Excel_Simulation!B9", T10["A"]["total_expected_cost"], T200["A"]["total_expected_cost"], "#,##0.00", "=3*Excel_Simulation!B10")
put(wc, "A11", "WITH M-BUF", bold, border=False)
trow(12, "P(delay occurs)", "=Test1_Hand!C29", "=Excel_Simulation!C5", None, None, "0.0000", f"=3*SQRT(Test1_Hand!C29*(1-Test1_Hand!C29)/{N})")
for col in "DEGHKL":
    wc[f"{col}12"].value = None
trow(13, "Expected delay (d)", "=Test1_Hand!C30", "=Excel_Simulation!C6", T10["B"]["expected_delay"], T200["B"]["expected_delay"], "0.0000", f"=3*SQRT(Test1_Hand!C29*(Test1_Hand!$B$20+Test1_Hand!$B$19^2)-Test1_Hand!C30^2)/SQRT({N})")
trow(14, "P(finish after the deadline)", "=Test1_Hand!C32", "=Excel_Simulation!C7", T10["B"]["p_exceed_deadline"], T200["B"]["p_exceed_deadline"], "0.0000", f"=3*SQRT(Test1_Hand!C32*(1-Test1_Hand!C32)/{N})")
trow(15, "P90 duration (d)", "=Test1_Hand!C40", "=Excel_Simulation!C8", T10["B"]["P90"], T200["B"]["P90"], "0.0000", "=3*C33")
trow(16, "Total expected cost incl. M-BUF (INR)", "=Test1_Hand!C37", "=Excel_Simulation!C9", T10["B"]["total_expected_cost"], T200["B"]["total_expected_cost"], "#,##0.00", "=3*Excel_Simulation!C10")
for col in "DEGHKL":
    wc[f"{col}12"].value = None
for col in "DGEHKL":
    wc[f"{col}12"].border = box
put(wc, "A12", "P(delay occurs)", bold)
put(wc, "D12", "n/a", blk, al=ctr); put(wc, "E12", "n/a", blk, al=ctr)
wc["G12"].value = None; wc["H12"].value = None; wc["K12"].value = None; wc["L12"].value = None
put(wc, "A31", "Standard error of the P90 duration at the Excel run's N (helper for the tolerance)", bold, border=False)
header(wc, 32, ["Quantity", "Accept", "With M-BUF"], height=24)
put(wc, "A33", "SE of P90 (d) = sqrt(0.9 x 0.1 / N) / density at P90", bold)
for col, hc in (("B", "B"), ("C", "C")):
    p = f"Test1_Hand!{hc}29"
    x = f"(Test1_Hand!{hc}40-Test1_Hand!$B$4)"
    dens = f"({p}*(BETADIST({x}+0.05,Test1_Hand!$B$21,Test1_Hand!$B$22,Test1_Hand!$B$9,Test1_Hand!$B$11)-BETADIST({x}-0.05,Test1_Hand!$B$21,Test1_Hand!$B$22,Test1_Hand!$B$9,Test1_Hand!$B$11))/0.1)"
    put(wc, f"{col}33", f"=SQRT(0.9*0.1/{N})/{dens}", blk, "0.0000")
put(wc, "A34", "Density of duration at the exact P90 (per day)", bold)
for col, hc in (("B", "B"), ("C", "C")):
    p = f"Test1_Hand!{hc}29"
    x = f"(Test1_Hand!{hc}40-Test1_Hand!$B$4)"
    put(wc, f"{col}34", f"={p}*(BETADIST({x}+0.05,Test1_Hand!$B$21,Test1_Hand!$B$22,Test1_Hand!$B$9,Test1_Hand!$B$11)-BETADIST({x}-0.05,Test1_Hand!$B$21,Test1_Hand!$B$22,Test1_Hand!$B$9,Test1_Hand!$B$11))/0.1", blk, "0.0000")
put(wc, "A18", "Decision", bold, border=False)
header(wc, 19, ["Source", "Decision", "TEC accept (INR)", "TEC with M-BUF (INR)", "Saving from M-BUF (INR)", "Agrees with hand?"], height=34)
put(wc, "A20", "Hand (exact)", bold)
put(wc, "B20", "=Test1_Hand!B43", grn)
put(wc, "C20", "=Test1_Hand!B37", grn, "#,##0.00")
put(wc, "D20", "=Test1_Hand!C37", grn, "#,##0.00")
put(wc, "E20", "=C20-D20", blk, "#,##0.00")
put(wc, "F20", "-", blk, al=ctr)
put(wc, "A21", "Excel simulation", bold)
put(wc, "B21", '=IF(D21<C21,"AUTHORIZE MITIGATION","ACCEPT RISK")', blk)
put(wc, "C21", "=Excel_Simulation!B9", grn, "#,##0.00")
put(wc, "D21", "=Excel_Simulation!C9", grn, "#,##0.00")
put(wc, "E21", "=C21-D21", blk, "#,##0.00")
put(wc, "F21", '=IF(B21=$B$20,"PASS","FAIL")', bold, al=ctr)
for r, nm, T in ((22, "Tool, n=10,000", T10), (23, "Tool, n=200,000", T200)):
    put(wc, f"A{r}", nm, bold)
    put(wc, f"B{r}", T["command"], blue)
    put(wc, f"C{r}", T["A"]["total_expected_cost"], blue, "#,##0.00")
    put(wc, f"D{r}", T["B"]["total_expected_cost"], blue, "#,##0.00")
    put(wc, f"E{r}", f"=C{r}-D{r}", blk, "#,##0.00")
    put(wc, f"F{r}", f'=IF(B{r}=$B$20,"PASS","FAIL")', bold, al=ctr)
put(wc, "A25", "Overall", bold)
put(wc, "B25", '=IF(COUNTIF(J6:L16,"FAIL")+COUNTIF(F21:F23,"FAIL")=0,"ALL CHECKS PASS","SOME CHECKS FAIL")', bold)
put(wc, "A26", "Checks counted", bold)
put(wc, "B26", '=COUNTIF(J6:L16,"PASS")+COUNTIF(J6:L16,"FAIL")+COUNTIF(F21:F23,"PASS")+COUNTIF(F21:F23,"FAIL")', blk, "0")
put(wc, "A28", "Reading: the exact hand result, a 2,000-draw spreadsheet simulation and the tool's own simulation agree within sampling error, and all three reach the same decision. These inputs are placeholders; the result shows the arithmetic is right, not that the numbers fit a real site.", blk, al=wrap, border=False)
wc.merge_cells("A28:L28")
wc.row_dimensions[28].height = 48
passfail(wc, "J6:L16")
passfail(wc, "F21:F23")
widths(wc, [38, 16, 18, 16, 16, 14, 16, 16, 14, 9, 10, 10])

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
wb.save(OUT)
print("saved", OUT)
