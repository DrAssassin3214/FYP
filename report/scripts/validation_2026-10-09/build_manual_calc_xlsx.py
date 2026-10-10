"""Build the manual-calculation workbook. Inputs come from the example case; tool values from the saved validation results."""
import json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule

REPO = Path(__file__).resolve().parents[3]
OUT = Path(sys.argv[1])
case = json.load(open(REPO / "examples/example_analysis_case.json", encoding="utf-8-sig"))
RES = json.load(open(REPO / "report/scripts/validation_2026-10-09/v1_results.json"))
R, CH = RES["results"], RES["checks"]

F = "Arial"
blue = Font(name=F, size=10, color="0000FF")
blk = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
grn = Font(name=F, size=10, color="008000")
title = Font(name=F, size=14, bold=True)
hdr_font = Font(name=F, size=10, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", start_color="1F3864")
inp_fill = PatternFill("solid", start_color="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()


def put(ws, ref, v, font=blk, fmt=None, fill=None, al=None, border=True):
    c = ws[ref]
    c.value = v
    c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if al: c.alignment = al
    if border: c.border = box
    return c


def header(ws, row, labels, col0=1, height=32):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=t)
        c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, ctr, box
    ws.row_dimensions[row].height = height


def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def passfail(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], font=Font(name=F, color="006100", bold=True), fill=PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], font=Font(name=F, color="9C0006", bold=True), fill=PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")))


# ------------------------------------------------------------------------------------------------ README
ws = wb.active
ws.title = "README"
ws["A1"].value, ws["A1"].font = "Manual-calculation check of the brick-masonry delay-risk tool", title
lines = [
    ("What this is", "Hand calculations, written as live Excel formulas, for the ILLUSTRATIVE example case (examples/example_analysis_case.json). Each tool value is the number printed by the tool in the validation run of 9 Oct 2026; the workbook recomputes the same quantity by hand and shows the difference."),
    ("Not evidence", "Every probability, delay range, duration and cost in this example is an Assumption placeholder, not site data. This workbook shows the arithmetic is right, not that the numbers describe a real site."),
    ("Colour legend", "Blue text on yellow = input you may change. Black = formula. Green = link to another sheet. Tool values (pasted from the validation run) are blue text on white."),
    ("Sheets", "Inputs: case numbers. Risk_Calc: PERT mean, variance, EMV, totals. Matrix: 5x5 class and level by hand. Mitigation_Calc: net benefit and break-even. Options: 23 options, exact total expected cost vs tool. Copula: occurrence-correlation check. All_Checks: all 242 checks of the run. Test_Summary: counts and findings."),
    ("Exact values", "Option costs on the Options sheet and the copula joint probabilities come from numerical convolution and a bivariate-normal integral (scripts in report/scripts/validation_2026-10-09/). They cannot be written as a cell formula, so they are pasted as values; every sum, rank and difference around them is a formula."),
    ("Changing inputs", "If you change a blue input on Inputs, the hand columns recompute but the pasted tool columns do not, so differences will appear. Re-run the tool and paste new tool values to compare again."),
    ("Source", "Case: examples/example_analysis_case.json. Tool and exact values: report/scripts/validation_2026-10-09/v1_results.json (run of 9 Oct 2026, seed 12345, n = 10,000 and 200,000). Formulas: Beta-PERT mean (a + lambda*m + b)/(lambda + 2) and variance (mean - a)(b - mean)/(lambda + 3); event EMV = p x mean delay x delay cost per day."),
]
r = 3
for k, v in lines:
    put(ws, f"A{r}", k, bold, al=wrap)
    put(ws, f"B{r}", v, blk, al=wrap)
    ws.row_dimensions[r].height = 62
    r += 1
widths(ws, [20, 120])

# ------------------------------------------------------------------------------------------------ Inputs
wi = wb.create_sheet("Inputs")
wi["A1"].value, wi["A1"].font = "Inputs (ILLUSTRATIVE example case; all values are Assumptions)", title
header(wi, 3, ["Quantity", "Value", "Unit", "Note"])
base = [
    ("Planned duration T0", case["activity"]["planned_duration_days"]["value"], "days", "activity.planned_duration_days"),
    ("Deadline", case["activity"]["deadline_days"]["value"], "days", "activity.deadline_days"),
    ("Delay cost Cd", case["cost"]["cost_per_delay_day"]["value"], "INR/day", "cost.cost_per_delay_day"),
    ("Liquidated damages LD", case["cost"]["ld_per_day_after_deadline"]["value"], "INR/day", "after the deadline"),
    ("PERT lambda", 4, "-", "standard Beta-PERT weight; the tool's default"),
]
for i, (k, v, u, n) in enumerate(base):
    rr = 4 + i
    put(wi, f"A{rr}", k, bold)
    put(wi, f"B{rr}", v, blue, "#,##0.00" if k == "PERT lambda" else "#,##0", inp_fill)
    put(wi, f"C{rr}", u)
    put(wi, f"D{rr}", n)
# named positions: T0=B4, DL=B5, CD=B6, LD=B7, LAM=B8
header(wi, 11, ["Risk id", "Name", "p (probability)", "a (min, days)", "m (mode, days)", "b (max, days)"])
RISKS = case["risks"]
for i, rk in enumerate(RISKS):
    rr = 12 + i
    put(wi, f"A{rr}", rk["id"], bold)
    put(wi, f"B{rr}", rk["name"])
    put(wi, f"C{rr}", rk["p"]["value"], blue, "0.00", inp_fill)
    put(wi, f"D{rr}", rk["delay"]["a"], blue, "0.0", inp_fill)
    put(wi, f"E{rr}", rk["delay"]["m"], blue, "0.0", inp_fill)
    put(wi, f"F{rr}", rk["delay"]["b"], blue, "0.0", inp_fill)
R0, R1 = 12, 12 + len(RISKS) - 1       # 12..18
M_HDR = R1 + 3                        # 21
header(wi, M_HDR, ["Response id", "Targets risk", "Cost (INR)", "p after", "Delay after a", "Delay after m", "Delay after b", "Secondary p", "Secondary a", "Secondary m", "Secondary b", "Feasible?"], height=40)
MITS = case["mitigations"]
for i, m in enumerate(MITS):
    rr = M_HDR + 1 + i
    da = m.get("delay_after")
    sr = (m.get("secondary_risks") or [None])[0]
    put(wi, f"A{rr}", m["id"], bold)
    put(wi, f"B{rr}", m["risk_id"])
    put(wi, f"C{rr}", m["cost"]["value"], blue, "#,##0", inp_fill)
    put(wi, f"D{rr}", m["p_after"]["value"] if m.get("p_after") else None, blue, "0.00", inp_fill)
    for col, key in zip("EFG", ("a", "m", "b")):
        put(wi, f"{col}{rr}", da[key] if da else None, blue, "0.0", inp_fill)
    put(wi, f"H{rr}", sr["p"]["value"] if sr else None, blue, "0.00", inp_fill)
    for col, key in zip("IJK", ("a", "m", "b")):
        put(wi, f"{col}{rr}", sr["delay"][key] if sr else None, blue, "0.0", inp_fill)
    put(wi, f"L{rr}", "No" if m.get("feasible") is False else "Yes", blue, None, inp_fill)
MR0, MR1 = M_HDR + 1, M_HDR + len(MITS)   # 22..27
wi.cell(row=MR1 + 2, column=1, value="Matrix settings").font = bold
put(wi, f"A{MR1+3}", "Probability class edges (lower-edge inclusive)", bold)
for j, e in enumerate((0.2, 0.4, 0.6, 0.8)):
    put(wi, f"{get_column_letter(2+j)}{MR1+3}", e, blue, "0.00", inp_fill)
put(wi, f"A{MR1+4}", "Impact edges (fraction of T0)", bold)
for j, e in enumerate(case["impact_bin_edges_fraction"]):
    put(wi, f"{get_column_letter(2+j)}{MR1+4}", e, blue, "0.00", inp_fill)
put(wi, f"A{MR1+5}", "Level cut-offs on score (Low <=, Moderate <=, High <=)", bold)
for j, e in enumerate((5, 10, 15)):
    put(wi, f"{get_column_letter(2+j)}{MR1+5}", e, blue, "0", inp_fill)
PE_ROW, IE_ROW, LV_ROW = MR1 + 3, MR1 + 4, MR1 + 5
widths(wi, [46, 38, 16, 14, 14, 14, 14, 12, 12, 12, 12, 11])

# ------------------------------------------------------------------------------------------------ Risk_Calc
wr = wb.create_sheet("Risk_Calc")
wr["A1"].value, wr["A1"].font = "Per-risk hand calculation vs tool", title
header(wr, 3, ["Risk id", "p", "Mean delay (d)\n=(a+lam*m+b)/(lam+2)", "Variance (d^2)\n=(mean-a)(b-mean)/(lam+3)", "EMV (INR)\n=p*mean*Cd", "p x mean (d)", "Var contribution\n=p(var+mean^2)-(p mean)^2", "1 - p",
                "Tool mean", "Tool EMV (INR)", "Mean diff", "EMV diff (INR)", "Result"], height=52)
V2 = {x["risk"]: x for x in R["V2"]}
for i, rk in enumerate(RISKS):
    rr = 4 + i
    ir = R0 + i
    put(wr, f"A{rr}", f"=Inputs!A{ir}", grn)
    put(wr, f"B{rr}", f"=Inputs!C{ir}", grn, "0.00")
    put(wr, f"C{rr}", f"=(Inputs!D{ir}+Inputs!$B$8*Inputs!E{ir}+Inputs!F{ir})/(Inputs!$B$8+2)", blk, "0.0000")
    put(wr, f"D{rr}", f"=(C{rr}-Inputs!D{ir})*(Inputs!F{ir}-C{rr})/(Inputs!$B$8+3)", blk, "0.000000")
    put(wr, f"E{rr}", f"=B{rr}*C{rr}*Inputs!$B$6", blk, "#,##0.00")
    put(wr, f"F{rr}", f"=B{rr}*C{rr}", blk, "0.00000")
    put(wr, f"G{rr}", f"=B{rr}*(D{rr}+C{rr}^2)-F{rr}^2", blk, "0.000000")
    put(wr, f"H{rr}", f"=1-B{rr}", blk, "0.00")
    put(wr, f"I{rr}", V2[rk["id"]]["tool_mean"], blue, "0.0000")
    put(wr, f"J{rr}", V2[rk["id"]]["tool_emv"], blue, "#,##0.00")
    put(wr, f"K{rr}", f"=C{rr}-I{rr}", blk, "0.0E+00")
    put(wr, f"L{rr}", f"=E{rr}-J{rr}", blk, "0.0E+00")
    put(wr, f"M{rr}", f'=IF(AND(ABS(K{rr})<=0.000001,ABS(L{rr})<=0.0001),"PASS","FAIL")', bold, al=ctr)
e0, e1 = 4, 3 + len(RISKS)    # 4..10
tr = e1 + 2                   # 12
put(wr, f"A{tr}", "Totals", bold)
tot = [
    ("Expected total delay E[D] (d)", f"=SUM(F{e0}:F{e1})", "0.00000", R["V2_totals"]["E_delay"], "0.00000"),
    ("Variance of total delay (d^2)", f"=SUM(G{e0}:G{e1})", "0.00000", R["V2_totals"]["Var"], "0.00000"),
    ("Standard deviation (d)", f"=SQRT(B{tr+2})", "0.00000", R["V2_totals"]["sd"], "0.00000"),
    ("Sum of event EMV (INR)", f"=SUM(E{e0}:E{e1})", "#,##0.00", R["V2_totals"]["EMV_sum"], "#,##0.00"),
    ("P(at least one delay)  =1-PRODUCT(1-p)", f"=1-PRODUCT(H{e0}:H{e1})", "0.000000", R["V2_totals"]["P_any"], "0.000000"),
    ("Planned duration + E[D] (d)", f"=Inputs!B4+B{tr+1}", "0.0000", None, None),
]
header(wr, tr + 0, ["Totals", "Hand (formula)", "Tool / exact", "Difference", "Result"], height=24)
for i, (k, f, fm, tv, tfm) in enumerate(tot):
    rr = tr + 1 + i
    put(wr, f"A{rr}", k, bold)
    put(wr, f"B{rr}", f, blk, fm)
    if tv is not None:
        put(wr, f"C{rr}", tv, blue, tfm)
        put(wr, f"D{rr}", f"=B{rr}-C{rr}", blk, "0.0E+00")
        put(wr, f"E{rr}", f'=IF(ABS(D{rr})<=0.0001,"PASS","FAIL")', bold, al=ctr)
# fix SD reference: SD row refers to variance row (tr+2)
wr[f"B{tr+3}"].value = f"=SQRT(B{tr+2})"
wr.cell(row=tr + 9, column=1, value="Note: Var contribution uses E[(I*D)^2] = p(var+mean^2) and the independent-occurrence assumption (no correlation), the case the tool runs by default. Tool values are from the validation run; the tool variance is the closed form used in the tool's own tests.").font = blk
passfail(wr, f"M{e0}:M{e1}")
passfail(wr, f"E{tr+1}:E{tr+6}")
widths(wr, [40, 15, 20, 24, 18, 14, 26, 10, 12, 16, 12, 14, 10])

# ------------------------------------------------------------------------------------------------ Matrix
wm = wb.create_sheet("Matrix")
wm["A1"].value, wm["A1"].font = "5 x 5 matrix class and level by hand vs tool", title
header(wm, 3, ["Risk id", "p", "Mean delay (d)", "Ratio = mean / T0", "Probability class\n=1+count(p>=edge)", "Impact class\n=1+count(ratio>=edge)", "Score = pc x ic", "Level (hand)", "Tool pc", "Tool ic", "Tool score", "Tool level", "Result"], height=52)
V5 = {x["risk"]: x for x in R["V5_rows"]}
for i, rk in enumerate(RISKS):
    rr = 4 + i
    put(wm, f"A{rr}", f"=Risk_Calc!A{4+i}", grn)
    put(wm, f"B{rr}", f"=Risk_Calc!B{4+i}", grn, "0.00")
    put(wm, f"C{rr}", f"=Risk_Calc!C{4+i}", grn, "0.0000")
    put(wm, f"D{rr}", f"=C{rr}/Inputs!$B$4", blk, "0.00000")
    put(wm, f"E{rr}", f"=1+(B{rr}>=Inputs!$B${PE_ROW}-0.000000000001)+(B{rr}>=Inputs!$C${PE_ROW}-0.000000000001)+(B{rr}>=Inputs!$D${PE_ROW}-0.000000000001)+(B{rr}>=Inputs!$E${PE_ROW}-0.000000000001)", blk, "0")
    put(wm, f"F{rr}", f"=1+(D{rr}>=Inputs!$B${IE_ROW}-0.000000000001)+(D{rr}>=Inputs!$C${IE_ROW}-0.000000000001)+(D{rr}>=Inputs!$D${IE_ROW}-0.000000000001)+(D{rr}>=Inputs!$E${IE_ROW}-0.000000000001)", blk, "0")
    put(wm, f"G{rr}", f"=E{rr}*F{rr}", blk, "0")
    put(wm, f"H{rr}", f'=IF(G{rr}<=Inputs!$B${LV_ROW},"Low",IF(G{rr}<=Inputs!$C${LV_ROW},"Moderate",IF(G{rr}<=Inputs!$D${LV_ROW},"High","Extreme")))', blk)
    t = V5[rk["id"]]
    put(wm, f"I{rr}", t["pc"], blue, "0")
    put(wm, f"J{rr}", t["ic"], blue, "0")
    put(wm, f"K{rr}", t["score"], blue, "0")
    put(wm, f"L{rr}", t["level"], blue)
    put(wm, f"M{rr}", f'=IF(AND(E{rr}=I{rr},F{rr}=J{rr},G{rr}=K{rr},H{rr}=L{rr}),"PASS","FAIL")', bold, al=ctr)
passfail(wm, f"M4:M{3+len(RISKS)}")
wm.cell(row=13, column=1, value="Exhaustive check in the validation run: 0 of 20,010 probability-class, 0 of 20,004 impact-class and 0 of 25 level mismatches (All_Checks, V5). Edges are lower-edge inclusive; the edge tolerance is 1e-12.").font = blk
widths(wm, [12, 8, 15, 16, 20, 20, 14, 14, 9, 9, 11, 12, 10])

# ------------------------------------------------------------------------------------------------ Mitigation_Calc
wc = wb.create_sheet("Mitigation_Calc")
wc["A1"].value, wc["A1"].font = "Net benefit of each response, by hand vs tool", title
header(wc, 3, ["Response", "Targets", "EMV before (INR)", "Mean delay after (d)", "p after", "EMV after (INR)", "Secondary EMV (INR)", "Cost (INR)", "Net benefit (INR)\n=before-after-secondary-cost", "Break-even reduction\n=cost/EMV before", "Pays?", "Feasible?",
                "Tool net (INR)", "Tool break-even", "Net diff", "Result"], height=60)
V4 = {x["mit"]: x for x in R["V4"]}
for i, m in enumerate(MITS):
    rr = 4 + i
    ir = MR0 + i
    put(wc, f"A{rr}", f"=Inputs!A{ir}", grn)
    put(wc, f"B{rr}", f"=Inputs!B{ir}", grn)
    put(wc, f"C{rr}", f"=INDEX(Risk_Calc!$E$4:$E${e1},MATCH(B{rr},Risk_Calc!$A$4:$A${e1},0))", grn, "#,##0.00")
    # mean delay after: uses delay_after if given, else original mean
    put(wc, f"D{rr}", f'=IF(ISNUMBER(Inputs!E{ir}),(Inputs!E{ir}+Inputs!$B$8*Inputs!F{ir}+Inputs!G{ir})/(Inputs!$B$8+2),INDEX(Risk_Calc!$C$4:$C${e1},MATCH(B{rr},Risk_Calc!$A$4:$A${e1},0)))', blk, "0.0000")
    put(wc, f"E{rr}", f'=IF(ISNUMBER(Inputs!D{ir}),Inputs!D{ir},INDEX(Risk_Calc!$B$4:$B${e1},MATCH(B{rr},Risk_Calc!$A$4:$A${e1},0)))', blk, "0.00")
    put(wc, f"F{rr}", f"=E{rr}*D{rr}*Inputs!$B$6", blk, "#,##0.00")
    put(wc, f"G{rr}", f'=IF(ISNUMBER(Inputs!H{ir}),Inputs!H{ir}*(Inputs!I{ir}+Inputs!$B$8*Inputs!J{ir}+Inputs!K{ir})/(Inputs!$B$8+2)*Inputs!$B$6,0)', blk, "#,##0.00")
    put(wc, f"H{rr}", f"=Inputs!C{ir}", grn, "#,##0")
    put(wc, f"I{rr}", f"=C{rr}-F{rr}-G{rr}-H{rr}", blk, "#,##0.00;(#,##0.00);-")
    put(wc, f"J{rr}", f"=H{rr}/C{rr}", blk, "0.0000")
    put(wc, f"K{rr}", f'=IF(I{rr}>0,"Yes","No")', blk, al=ctr)
    put(wc, f"L{rr}", f"=Inputs!L{ir}", grn, al=ctr)
    t = V4[m["id"]]
    put(wc, f"M{rr}", t["tool_net"], blue, "#,##0.00;(#,##0.00);-")
    put(wc, f"N{rr}", t["tool_required_reduction"], blue, "0.0000")
    put(wc, f"O{rr}", f"=I{rr}-M{rr}", blk, "0.0E+00")
    put(wc, f"P{rr}", f'=IF(AND(ABS(O{rr})<=0.01,ABS(J{rr}-N{rr})<=0.000001),"PASS","FAIL")', bold, al=ctr)
passfail(wc, f"P4:P{3+len(MITS)}")
wc.cell(row=11, column=1, value="Reading: a response pays when the EMV it removes exceeds its cost plus any secondary risk it creates. M-MECH is marked not feasible in the case, so the tool keeps it out of the options even though its arithmetic net benefit is shown here.").font = blk
widths(wc, [11, 10, 16, 16, 9, 16, 16, 12, 26, 20, 8, 10, 15, 14, 12, 10])

# ------------------------------------------------------------------------------------------------ Options
wo = wb.create_sheet("Options")
wo["A1"].value, wo["A1"].font = "All 23 options: exact total expected cost (TEC) vs tool", title
wo["A2"].value = "Residual expected cost and exact P90 come from numerical convolution of the mixture distributions (pasted values, blue). TEC, ranks, z-scores and the decision are formulas."
wo["A2"].font = blk
header(wo, 4, ["Option", "Mitigation cost (INR)", "Exact residual E[cost] (INR)", "Exact TEC (INR)\n=mit cost + residual", "Rank among feasible\n(1 = lowest TEC)", "Tool TEC n=10,000 (INR)", "Tool TEC n=200,000 (INR)", "Standard error n=200k (INR)", "z-score n=200k\n=(tool-exact)/SE", "Exact P90 (d)", "Tool P90 n=200k (d)", "Exact P(finish > deadline)", "Tool P(exceed) n=200k", "Feasible?", "|z| <= 4?"], height=64)
HO = R["V3_hand_options"]
rows = R["V3_rows"]
for i, rw in enumerate(rows):
    rr = 5 + i
    put(wo, f"A{rr}", rw["option"], bold)
    put(wo, f"B{rr}", rw["mit_cost"], blue, "#,##0")
    put(wo, f"C{rr}", HO[rw["option"]]["E_cost"], blue, "#,##0.00")
    put(wo, f"D{rr}", f"=B{rr}+C{rr}", blk, "#,##0.00")
    put(wo, f"F{rr}", rw["tool_TEC_10k"], blue, "#,##0.00")
    put(wo, f"G{rr}", rw["tool_TEC_200k"], blue, "#,##0.00")
    put(wo, f"H{rr}", rw["se200"], blue, "#,##0.00")
    put(wo, f"I{rr}", f"=(G{rr}-D{rr})/H{rr}", blk, "0.00")
    put(wo, f"J{rr}", rw["hand_P90"], blue, "0.000")
    put(wo, f"K{rr}", rw["tool_P90_200k"], blue, "0.000")
    put(wo, f"L{rr}", rw["hand_pexc"], blue, "0.0000")
    put(wo, f"M{rr}", rw["tool_pexc_200k"], blue, "0.0000")
    put(wo, f"N{rr}", "Yes" if rw["feasible"] else "No", blue, al=ctr)
    put(wo, f"O{rr}", f'=IF(ABS(I{rr})<=4,"PASS","FAIL")', bold, al=ctr)
o0, o1 = 5, 4 + len(rows)
for rr in range(o0, o1 + 1):
    put(wo, f"E{rr}", f'=IF(N{rr}="Yes",COUNTIFS($N${o0}:$N${o1},"Yes",$D${o0}:$D${o1},"<"&D{rr})+1,"-")', blk, "0", al=ctr)
passfail(wo, f"O{o0}:O{o1}")
d = o1 + 2
put(wo, f"A{d}", "Hand-best feasible option", bold)
put(wo, f"B{d}", f'=INDEX($A${o0}:$A${o1},MATCH(1,$E${o0}:$E${o1},0))', blk)
put(wo, f"A{d+1}", "Its exact TEC (INR)", bold)
put(wo, f"B{d+1}", f'=INDEX($D${o0}:$D${o1},MATCH(1,$E${o0}:$E${o1},0))', blk, "#,##0.00")
put(wo, f"A{d+2}", "Tool's chosen option, n=10,000", bold)
put(wo, f"B{d+2}", R["V3_decision"]["tool10k"], blue)
put(wo, f"A{d+3}", "Tool's chosen option, n=200,000", bold)
put(wo, f"B{d+3}", R["V3_decision"]["tool200k"], blue)
put(wo, f"A{d+4}", "Decision agrees?", bold)
put(wo, f"B{d+4}", f'=IF(AND(B{d}=B{d+2},B{d}=B{d+3}),"PASS","FAIL")', bold, al=ctr)
put(wo, f"A{d+5}", "Tool command at n=10,000", bold)
put(wo, f"B{d+5}", R["V3_decision"]["command10k"], blue)
put(wo, f"A{d+6}", "Gap: 2nd best minus best TEC (INR)", bold)
put(wo, f"B{d+6}", f'=SUMPRODUCT(($E${o0}:$E${o1}=2)*$D${o0}:$D${o1})-SUMPRODUCT(($E${o0}:$E${o1}=1)*$D${o0}:$D${o1})', blk, "#,##0.0")
passfail(wo, f"B{d+4}")
wo.cell(row=d + 8, column=1, value="A gap of a few hundred INR between first and second option is small against a cost scale of about 44,000 INR: treat the choice between them as not firm. Largest |z| at n=200,000 is 2.97 for one of 23 options, consistent with chance (see Test_Summary).").font = blk
widths(wo, [36, 16, 20, 18, 16, 18, 18, 16, 16, 12, 14, 16, 14, 10, 10])
wo.freeze_panes = "B5"

# ------------------------------------------------------------------------------------------------ Copula
wp = wb.create_sheet("Copula")
wp["A1"].value, wp["A1"].font = "Occurrence correlation (Gaussian copula): joint probability vs bivariate-normal value", title
header(wp, 3, ["rho", "Exact P(A and B)\n(bivariate normal)", "Tool P(A and B)", "Difference", "Tool marginal A", "Tool marginal B", "Result"], height=40)
for i, x in enumerate(R["V10"]):
    rr = 4 + i
    put(wp, f"A{rr}", x["rho"], blue, "0.00")
    put(wp, f"B{rr}", x["hand_joint"], blue, "0.0000")
    put(wp, f"C{rr}", x["tool_joint"], blue, "0.0000")
    put(wp, f"D{rr}", f"=C{rr}-B{rr}", blk, "0.0000")
    put(wp, f"E{rr}", x["marg_A"], blue, "0.0000")
    put(wp, f"F{rr}", x["marg_B"], blue, "0.0000")
    put(wp, f"G{rr}", f'=IF(ABS(D{rr})<=0.001,"PASS","FAIL")', bold, al=ctr)
passfail(wp, "G4:G7")
wp.cell(row=9, column=1, value="Independent case (rho = 0) must equal pA x pB by hand:").font = bold
put(wp, "A10", "pA x pB", bold)
put(wp, "B10", "=0.4*0.3", blk, "0.0000")
put(wp, "C10", "Marginals set to 0.40 and 0.30 in the test.", blk, border=False)
widths(wp, [14, 22, 18, 14, 16, 16, 10])

# ------------------------------------------------------------------------------------------------ All_Checks
wa = wb.create_sheet("All_Checks")
wa["A1"].value, wa["A1"].font = "All 242 checks of the hand-calculation run (9 Oct 2026, code before fixes F1-F6; F6 case in row for V11)", title
header(wa, 3, ["Check id", "Description", "Hand / exact", "Tool", "|Difference|\n(formula)", "Tolerance", "Result\n(formula)", "Result in run", "Note"], height=40)
for i, c in enumerate(CH):
    rr = 4 + i
    put(wa, f"A{rr}", c["id"])
    put(wa, f"B{rr}", c["desc"])
    put(wa, f"C{rr}", c["hand"], blue, "General")
    put(wa, f"D{rr}", c["tool"], blue, "General")
    put(wa, f"E{rr}", f"=ABS(C{rr}-D{rr})", blk, "0.0E+00")
    put(wa, f"F{rr}", c["tol"], blue, "0.0E+00")
    put(wa, f"G{rr}", f'=IF(E{rr}<=F{rr},"PASS","FAIL")', bold, al=ctr)
    put(wa, f"H{rr}", "PASS" if c["pass"] else "FAIL", blue, al=ctr)
    put(wa, f"I{rr}", c["note"] or "")
a0, a1 = 4, 3 + len(CH)
passfail(wa, f"G{a0}:H{a1}")
wa.freeze_panes = "C4"
wa.auto_filter.ref = f"A3:I{a1}"
widths(wa, [26, 72, 16, 16, 14, 12, 10, 10, 30])

# ------------------------------------------------------------------------------------------------ Test_Summary
wt = wb.create_sheet("Test_Summary")
wt["A1"].value, wt["A1"].font = "Test summary, 9 October 2026", title
header(wt, 3, ["Test set", "Checks", "Passed before fixes", "Passed after fixes", "Failed after", "Cause of remaining failures"], height=36)
summ = [
    ("Hand and exact calculation", 242, 239, 240, "Two P90 differences of 0.049 and 0.051 d at n=200,000: sampling noise"),
    ("CLI, API, front end, packaging flows", 86, 79, 85, "One wrong expectation in the test script (run on a case without a cost model exits 0 by design)"),
    ("Edge cases, wrong-shape inputs, performance", 85, 68, 83, "Two wrong expectations in the test script (lambda = 0 is correctly rejected; truncation notice is in the report)"),
]
for i, (k, n, b, a, why) in enumerate(summ):
    rr = 4 + i
    put(wt, f"A{rr}", k, bold)
    put(wt, f"B{rr}", n, blue, "#,##0")
    put(wt, f"C{rr}", b, blue, "#,##0")
    put(wt, f"D{rr}", a, blue, "#,##0")
    put(wt, f"E{rr}", f"=B{rr}-D{rr}", blk, "0")
    put(wt, f"F{rr}", why, al=wrap)
put(wt, "A7", "Total", bold)
for col in "BCD":
    put(wt, f"{col}7", f"=SUM({col}4:{col}6)", bold, "#,##0")
put(wt, "E7", "=B7-D7", bold, "0")
put(wt, "F7", "", blk)
header(wt, 9, ["Other test sets", "Size", "Before fixes", "After fixes", "", "Note"], height=24)
oth = [
    ("Fuzzing, mutated inputs", 7700, "12 unhandled errors", "0 unhandled", "4,000 register + 3,000 analysis-case + 700 run_analysis mutations"),
    ("Invariants on random valid cases", 140, "140 held", "140 held", "No valid case rejected"),
    ("Repository test suite", 557, "522 passed, 3 skipped (of 525)", "554 passed, 3 skipped", "32 new regression tests in tests/test_validation_fixes.py; 3 skipped need a database not built here"),
]
for i, (k, n, b, a, note) in enumerate(oth):
    rr = 10 + i
    put(wt, f"A{rr}", k, bold)
    put(wt, f"B{rr}", n, blue, "#,##0")
    put(wt, f"C{rr}", b, blue)
    put(wt, f"D{rr}", a, blue)
    put(wt, f"E{rr}", "")
    put(wt, f"F{rr}", note, al=wrap)
header(wt, 15, ["Finding", "Severity", "Where", "What happened", "Fix", "Status"], height=24)
finds = [
    ("F1", "Medium", "app/cli.py, analyze", "Bad case file ended in a traceback (6 of 6 file types)", "Same loader and error handling as run", "Fixed, tested"),
    ("F2", "Medium", "app/analysis.py", "Wrong-type containers crashed (14 of 19 shapes)", "Type-check each container, report a case problem", "Fixed, tested"),
    ("F3", "Low", "app/analysis.py, decision.py:96", "Mixed-type evidence ids failed when sorted", "Convert ids to text", "Fixed, tested"),
    ("F4", "Low", "simulation.py, convergence()", "Empty checkpoints for n < 10; n = 1 crashed", "At least one draw per checkpoint", "Fixed, tested"),
    ("F5", "Low", "simulation.py, simulate()", "Two ids with equal CRC32 share a random stream", "Refuse and ask for a rename (hash unchanged)", "Fixed, tested"),
    ("F6", "Low", "examples/example_analysis_case.json", "File differed from the generator function", "Regenerated; test compares them", "Fixed, tested"),
    ("F7", "Info", "app/engine/options.py", "Automatic search stops at 3 responses and 40 options", "Document in the GUI", "Open"),
]
for i, row in enumerate(finds):
    for j, v in enumerate(row):
        put(wt, f"{get_column_letter(1+j)}{16+i}", v, bold if j == 0 else blk, al=wrap)
wt.cell(row=24, column=1, value="Not verified: the PySide6 window and Windows executables (PySide6 not installable here), the browser UI operated by a person, LLM-assisted risk suggestions (no key), line coverage.").font = bold
wt.cell(row=25, column=1, value="Source: report/scripts/validation_2026-10-09/ (first-run and after_fixes outputs). Counts typed from those outputs.").font = blk
widths(wt, [44, 16, 34, 52, 44, 52])

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
wb.save(OUT)
print("saved", OUT)
